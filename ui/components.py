"""
Reusable Streamlit UI components for the Recon Engine.
Extracted and improved from app_starter.py utility functions.
"""

import os
import re
import uuid
import shutil
import streamlit as st
import pandas as pd
import openpyxl
from pathlib import Path

from core.file_utils import get_sheet_names_for_file


# =============================================================================
# FILE HELPERS
# =============================================================================

def get_files():
    """Get list of registered files (with blank first option for selectboxes)."""
    return [""] + list(st.session_state.get("uploaded_file_paths", {}).keys())


def get_file_path(filename):
    """Get absolute path for a registered file label."""
    return st.session_state.get("uploaded_file_paths", {}).get(filename, "")


def register_file_path(path, display_name=None, overwrite=True):
    """
    Register an existing local file path (no copy).
    Enables true in-place editing on the original file.
    """
    p = Path(path).expanduser()
    if not p.is_absolute():
        p = p.resolve()
    if not p.exists():
        raise FileNotFoundError(f"File does not exist: {p}")

    name = os.path.basename(str(p)) if not display_name else display_name
    st.session_state.uploaded_file_paths[name] = str(p)
    return str(p)


def register_uploaded_file(uploaded_file, workspace_dir="workspace"):
    """
    Save an uploaded file (from st.file_uploader) to the workspace directory
    and register it in session state.
    """
    ws = Path(workspace_dir)
    ws.mkdir(parents=True, exist_ok=True)

    dest = ws / uploaded_file.name
    with open(str(dest), "wb") as f:
        f.write(uploaded_file.getbuffer())

    abs_path = str(dest.resolve())
    st.session_state.uploaded_file_paths[uploaded_file.name] = abs_path
    return abs_path


# =============================================================================
# FILE TYPE DETECTION
# =============================================================================

def _is_excel_label(file_label):
    path = get_file_path(file_label) if file_label else ""
    lower = (path or "").lower()
    return lower.endswith(".xlsx") or lower.endswith(".xls")


def _is_csv_label(file_label):
    path = get_file_path(file_label) if file_label else ""
    lower = (path or "").lower()
    return lower.endswith(".csv") or lower.endswith(".txt")


# =============================================================================
# SHEET / COLUMN SELECTORS
# =============================================================================

def get_sheets_for_file(file_label):
    """Return available sheet names for a registered file."""
    if not file_label:
        return []
    path = get_file_path(file_label)
    if not path:
        return []
    lower = path.lower()
    if lower.endswith(".csv") or lower.endswith(".txt"):
        return []
    return get_sheet_names_for_file(path) or []


def sheet_selectbox(label, file_label, key):
    """
    Smart sheet selector dropdown driven by the selected file.
    Auto-selects if only one sheet exists.
    """
    if not file_label:
        st.selectbox(label, options=[""], index=0, key=f"{key}__no_file", disabled=True)
        return ""

    path = get_file_path(file_label)
    if not path:
        st.selectbox(label, options=[""], index=0, key=f"{key}__no_path", disabled=True)
        return ""

    lower = path.lower()
    if lower.endswith(".csv") or lower.endswith(".txt"):
        st.selectbox(label, options=["(CSV - no sheets)"], index=0, key=f"{key}__csv", disabled=True)
        return ""

    sheets = get_sheet_names_for_file(path) or []
    if not sheets:
        st.selectbox(label, options=["(No sheets detected)"], index=0, key=f"{key}__none", disabled=True)
        return ""

    # Auto-select if only one sheet
    if len(sheets) == 1:
        st.selectbox(label, options=sheets, index=0, key=key, disabled=False)
        return sheets[0]

    return st.selectbox(label, options=[""] + sheets, index=0, key=key)


def get_headers_for_file_sheet(file_label, sheet_name, header_row):
    """
    Returns column header names from a file/sheet at a specific header row.
    Works for both Excel and CSV files.
    """
    if not file_label:
        return []
    path = get_file_path(file_label)
    if not path:
        return []
    try:
        headers = []

        if _is_excel_label(file_label):
            if not sheet_name:
                return []
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
            for c in range(1, ws.max_column + 1):
                v = ws.cell(row=int(header_row), column=c).value
                if v is None:
                    continue
                s = str(v).strip()
                if not s:
                    continue
                headers.append(s)
            wb.close()

        elif _is_csv_label(file_label):
            n = max(1, int(header_row))
            dfh = pd.read_csv(path, header=None, nrows=n, dtype=str, keep_default_na=False)
            if dfh.shape[0] < n:
                return []
            row = dfh.iloc[n - 1].tolist()
            for v in row:
                s = str(v).strip() if v is not None else ""
                if not s:
                    continue
                headers.append(s)
        else:
            return []

        # De-dupe preserving order
        seen = set()
        out = []
        for h in headers:
            if h in seen:
                continue
            seen.add(h)
            out.append(h)
        return out
    except Exception:
        return []


def column_selectbox(label, file_label, sheet_name, header_row, key):
    """Column selector dropdown populated from file headers."""
    if not file_label or not sheet_name:
        st.selectbox(label, options=[""], index=0, key=f"{key}__no_fs", disabled=True)
        return ""
    headers = get_headers_for_file_sheet(file_label, sheet_name, header_row)
    if not headers:
        st.selectbox(label, options=["(No headers detected)"], index=0, key=f"{key}__no_hdrs", disabled=True)
        return ""
    return st.selectbox(label, options=[""] + headers, index=0, key=key)


def smart_column_selector(label, file_label, sheet_name, header_row, key, hint=None):
    """
    Enhanced column selector with auto-detection hints.
    If hint is provided (e.g., "amount|value|balance"), auto-selects the best match.
    """
    if not file_label or not sheet_name:
        st.selectbox(label, options=[""], index=0, key=f"{key}__no_fs", disabled=True)
        return ""
    headers = get_headers_for_file_sheet(file_label, sheet_name, header_row)
    if not headers:
        st.selectbox(label, options=["(No headers detected)"], index=0, key=f"{key}__no_hdrs", disabled=True)
        return ""

    default_idx = 0
    if hint and headers:
        patterns = hint.split("|")
        for i, h in enumerate(headers):
            if any(p.lower() in h.lower() for p in patterns):
                default_idx = i + 1  # +1 for blank option
                break

    return st.selectbox(label, [""] + headers, index=default_idx, key=key)


# =============================================================================
# INPUT HELPERS
# =============================================================================

def export_options_ui(key_prefix):
    """Per-step export option UI. Returns {"enabled": bool, "path": str}."""
    with st.expander("Export as File (optional)", expanded=False):
        enabled = st.checkbox("Export output after this step", value=False,
                              key=f"{key_prefix}_export_enabled")
        default_val = st.session_state.get("default_export_path", "") or ""
        path = st.text_input(
            "Export folder or file path",
            value=default_val,
            key=f"{key_prefix}_export_path",
            disabled=not enabled,
            help="Folder: /Users/you/Exports/  |  File: /Users/you/Exports/output.xlsx",
        )
        if enabled and not path.strip():
            st.warning("Export is enabled but no path is set.")
        return {"enabled": bool(enabled), "path": str(path or "")}


def end_row_input(label, key_prefix, default_last=True, default_number=10):
    """End row input with 'last' or specific number option."""
    mode_default = "last" if default_last else "custom"
    mode = st.selectbox(
        label,
        options=["last", "custom"],
        key=f"{key_prefix}_mode",
        index=0 if mode_default == "last" else 1,
    )
    if mode == "last":
        return "last"
    n = st.number_input("End row (number)", min_value=1, value=int(default_number),
                        key=f"{key_prefix}_num")
    return str(int(n))


def end_index_input(label, key_prefix, default_last=False, default_number=1):
    """End index input for insert/delete operations."""
    mode_default = "last" if default_last else "custom"
    mode = st.selectbox(
        label,
        options=["custom", "last"],
        key=f"{key_prefix}_mode",
        index=1 if mode_default == "last" else 0,
    )
    if mode == "last":
        return "last"
    n = st.number_input("End index (number)", min_value=1, value=int(default_number),
                        key=f"{key_prefix}_num")
    return str(int(n))


# =============================================================================
# STEP MANAGEMENT
# =============================================================================

def add_step(step_type, name, config):
    """Add a step to the workflow with validation."""
    # Validate file selections for steps that need them
    if step_type not in ("add_files",) and isinstance(config, dict):
        if "file" in config and not config.get("file"):
            st.warning("Please select a file first.")
            return
        for k in ("src_file", "tgt_file", "lookup_file"):
            if k in config and not config.get(k):
                st.warning(f"Please select {k.replace('_', ' ')}.")
                return
        # Enforce sheet selection for Excel files
        if "file" in config and "sheet" in config:
            if _is_excel_label(config.get("file")) and not config.get("sheet"):
                st.warning("Please select a sheet for the chosen Excel file.")
                return
        # SUMIFS source sheet
        if step_type == "sumifs":
            if _is_excel_label(config.get("src_workbook")) and not config.get("src_sheet"):
                st.warning("Please select a source sheet for the SUMIFS source file.")
                return
        # Multi-file sheet validations
        for file_key, sheet_key, msg in (
            ("lookup_file", "lookup_sheet", "Please select a sheet for the lookup file."),
            ("src_file", "src_sheet", "Please select a sheet for the source file."),
            ("tgt_file", "tgt_sheet", "Please select a sheet for the target file."),
        ):
            if file_key in config and sheet_key in config:
                if _is_excel_label(config.get(file_key)) and not config.get(sheet_key):
                    st.warning(msg)
                    return

        # Explicit VLOOKUP validations
        if step_type == "vlookup" and config.get("mode") == "explicit":
            required = [
                ("lookup_value_file", "Lookup Value file"),
                ("lookup_value_sheet", "Lookup Value sheet"),
                ("lookup_value_column", "Lookup Value column"),
                ("search_file", "Search Value file"),
                ("search_sheet", "Search Value sheet"),
                ("search_column", "Search Value column"),
                ("return_file", "Return Value file"),
                ("return_sheet", "Return Value sheet"),
                ("return_column", "Return Value column"),
                ("return_to_file", "Return To file"),
                ("return_to_sheet", "Return To sheet"),
                ("return_to_column", "Return To column"),
            ]
            for k, label in required:
                if not config.get(k):
                    st.warning(f"Please select {label}.")
                    return

        if step_type == "sort":
            if isinstance(config.get("sort_cols"), list) and not config.get("sort_cols"):
                st.warning("Please add at least one sort column.")
                return

    # Check edit mode
    editing_idx = st.session_state.get("editing_step_idx")
    editing_data = st.session_state.get("editing_step_data")

    if editing_idx is not None and editing_data is not None:
        existing_step = st.session_state.workflow_steps[editing_idx]
        existing_step["type"] = step_type
        existing_step["name"] = name
        existing_step["config"] = config
        st.session_state.workflow_steps[editing_idx] = existing_step

        st.session_state.editing_step_idx = None
        st.session_state.editing_step_data = None
        st.session_state.pop("adding_step", None)
        st.session_state["_last_added_step_name"] = f"(Updated) {name}"
        return

    # Normal: create new step
    step = {
        "id": uuid.uuid4().hex[:12],
        "type": step_type,
        "name": name,
        "config": config,
        "depends_on": [],
        "lane": "",
    }
    insert_at = st.session_state.get("insert_at", None)
    if isinstance(insert_at, int):
        idx = max(0, min(int(insert_at), len(st.session_state.workflow_steps)))
        st.session_state.workflow_steps.insert(idx, step)
        st.session_state.insert_at = None
        st.session_state._workflow_notice = ""
    else:
        st.session_state.workflow_steps.append(step)

    st.session_state["_last_added_step_name"] = name
    try:
        st.session_state.pop("adding_step", None)
    except Exception:
        pass


# =============================================================================
# EXPORT HELPERS
# =============================================================================

def _safe_slug(s, max_len=60):
    s = (s or "").strip().lower()
    s = re.sub(r"[^a-z0-9]+", "_", s)
    s = re.sub(r"_+", "_", s).strip("_")
    if not s:
        return "step"
    return s[:max_len]


def _determine_modified_file_labels(step_type, cfg):
    """Determine which file labels are modified by a step."""
    if step_type in ("add_files", "data_input"):
        return []
    if step_type in ("import", "append", "header_map"):
        return [cfg.get("tgt_file", "")]
    if step_type == "vlookup" and (cfg.get("mode") == "explicit" or cfg.get("lookup_value_file")):
        return [cfg.get("return_to_file", "")]
    return [cfg.get("file", "")]


def _build_export_path(user_path, source_path, step_num, step_type):
    """Build export file path from user settings."""
    up = (user_path or "").strip()
    if not up:
        up = (st.session_state.get("default_export_path") or "").strip()
    if not up:
        return ""

    src = Path(source_path)
    src_stem = src.stem
    suffix = f"__step{int(step_num):03d}_{_safe_slug(step_type)}"

    p = Path(up).expanduser()
    looks_like_file = p.suffix.lower() in (".xlsx", ".xls", ".csv", ".txt")

    if looks_like_file:
        base_dir = p.parent
        base_stem = p.stem
        out_stem = f"{base_stem}{suffix}"
    else:
        base_dir = p
        out_stem = f"{src_stem}{suffix}"

    out_path = (base_dir / f"{out_stem}.xlsx").resolve()
    return str(out_path)


def _export_to_xlsx(source_path, export_path):
    """Export a file to xlsx format at the given path."""
    src = Path(source_path)
    dst = Path(export_path)
    dst.parent.mkdir(parents=True, exist_ok=True)

    lower = src.suffix.lower()
    if lower in (".xlsx", ".xls"):
        shutil.copy2(str(src), str(dst))
        return
    if lower in (".csv", ".txt"):
        df = pd.read_csv(str(src))
        with pd.ExcelWriter(str(dst), engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Sheet1")
        return
    shutil.copy2(str(src), str(dst))


def prepare_non_destructive_export(step_num, step):
    """
    When export is enabled, copy original file to export path BEFORE the operation,
    so the operation runs on the copy, not the original.
    Returns (swapped_paths, original_mappings) or (None, []).
    """
    cfg = step.get("config", {}) or {}
    export_cfg = cfg.get("export") or {}
    if not isinstance(export_cfg, dict) or not export_cfg.get("enabled"):
        return None, []

    step_type = (step.get("type") or "").strip().lower()
    user_path = export_cfg.get("path") or export_cfg.get("dir") or ""

    labels = [x for x in _determine_modified_file_labels(step_type, cfg) if x]
    if not labels:
        return None, []

    swapped_paths = {}
    original_mappings = []

    for label in labels:
        src_path = get_file_path(label)
        if not src_path:
            continue

        out_path = _build_export_path(user_path, src_path, step_num, step_type)
        if not out_path:
            st.warning("Export enabled but no export path provided.")
            return None, []

        try:
            _export_to_xlsx(src_path, out_path)
            original_mappings.append((label, src_path))
            st.session_state.uploaded_file_paths[label] = str(Path(out_path).resolve())
            swapped_paths[label] = out_path
        except Exception as e:
            st.warning(f"Export copy failed for '{label}': {e}")
            for lbl, orig_path in original_mappings:
                st.session_state.uploaded_file_paths[lbl] = orig_path
            return None, []

    return swapped_paths, original_mappings


def finalize_non_destructive_export(step_num, swapped_paths, original_mappings):
    """Restore original file registrations and register export copies as new entries."""
    for label, orig_path in original_mappings:
        st.session_state.uploaded_file_paths[label] = orig_path

    for label, export_path in swapped_paths.items():
        export_name = os.path.basename(export_path)
        st.session_state.uploaded_file_paths[export_name] = str(Path(export_path).resolve())
        st.info(f"Step {step_num} exported (original unchanged): {export_path}")


# =============================================================================
# SESSION STATE HELPERS
# =============================================================================

def ensure_step_ids():
    """Ensure every workflow step has a stable ID."""
    steps = st.session_state.get("workflow_steps", []) or []
    changed = False
    for s in steps:
        if not isinstance(s, dict):
            continue
        if not s.get("id"):
            s["id"] = uuid.uuid4().hex[:12]
            changed = True
        if "depends_on" not in s:
            s["depends_on"] = []
            changed = True
        if "lane" not in s:
            s["lane"] = ""
            changed = True
    if changed:
        st.session_state.workflow_steps = steps


def init_full_session_state():
    """Initialize all session state variables for the multi-page app."""
    if "uploaded_file_paths" not in st.session_state or not isinstance(st.session_state.uploaded_file_paths, dict):
        st.session_state.uploaded_file_paths = {}
    if "workflow_steps" not in st.session_state:
        st.session_state.workflow_steps = []
    if "workflows" not in st.session_state:
        st.session_state.workflows = {}
    if "main_file_select" not in st.session_state:
        st.session_state.main_file_select = ""
    if "main_sheet_name" not in st.session_state:
        st.session_state.main_sheet_name = ""
    if "insert_at" not in st.session_state:
        st.session_state.insert_at = None
    if "editing_step_idx" not in st.session_state:
        st.session_state.editing_step_idx = None
    if "editing_step_data" not in st.session_state:
        st.session_state.editing_step_data = None
    if "_workflow_notice" not in st.session_state:
        st.session_state._workflow_notice = ""
    if "default_export_path" not in st.session_state:
        st.session_state.default_export_path = ""
    if "current_workflow_name" not in st.session_state:
        st.session_state.current_workflow_name = ""
