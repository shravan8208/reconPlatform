"""
Workflow execution engine.
Extracted from app_starter.py - handles sequential and DAG-based parallel execution.
"""

import os
import concurrent.futures as cf
import streamlit as st

from core import wb_cache
from operations import (
    run_formula_step, run_convert_to_values_step, run_formula_broadcast_step, run_copy_paste_step,
    run_replace_single, run_forward_fill_step, run_insert_delete_step,
    run_delete_columns_step, run_clear_columns_data_step,
    run_vlookup_step, run_vlookup_explicit_step, run_advance_vlookup_step,
    run_filter_step, run_sort_step,
    run_import_step, run_append_step, run_header_mapping_step,
    run_conditional_write_step, run_sumifs_step,
    run_merge_cells_step, run_data_input_step, run_write_cell_step,
    run_input_source_step, run_normalize_import_step,
    run_delete_by_condition_step, run_advanced_delete_step, run_folder_summary_step,
    run_format_step,
    run_pivot_step,
    run_split_export_step,
    run_sheet_updater_step,
    run_unpivot_step,
    run_delete_sheets_step,
)
from ui.components import (
    get_file_path, register_file_path, ensure_step_ids,
    _determine_modified_file_labels,
    prepare_non_destructive_export, finalize_non_destructive_export,
)
from openpyxl.utils import get_column_letter


def _resolve_column(file_path: str, sheet: str, cfg: dict) -> str:
    """
    Return the column letter for a formula step.
    If cfg['column_is_header'] is True, scan the header row of the sheet
    to find the matching header and return its letter. Otherwise return cfg['column'] as-is.
    """
    col = cfg.get("column", "A")
    if not cfg.get("column_is_header"):
        return col
    try:
        wb = wb_cache.load(file_path)
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        header_row = int(cfg.get("column_header_row", 1) or 1)
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=header_row, column=c).value
            if v is not None and str(v).strip() == str(col).strip():
                return get_column_letter(c)
    except Exception:
        pass
    return col  # fall back to whatever was stored (may be a header name string)


def execute_step(step):
    """Execute a single workflow step using session state for file resolution."""
    step_type = (step.get("type") or "").lower().strip()
    cfg = step.get("config", {}) or {}

    def _registered_labels():
        m = st.session_state.get("uploaded_file_paths")
        if not isinstance(m, dict):
            return []
        return sorted([str(k) for k in m.keys()])

    def _resolve(label, field):
        if not label:
            return "", f"Missing file for '{field}'."
        path = get_file_path(label)
        if path:
            return path, None
        reg = ", ".join([f"'{x}'" for x in _registered_labels()[:10]])
        return "", f"File not found: '{label}' for '{field}'. Registered: [{reg}]"

    # add_files step
    if step_type == "add_files":
        files = cfg.get("files", []) or []
        if not files:
            return False, "No files provided"
        registered = 0
        for item in files:
            name = item.get("name")
            path = item.get("path")
            if not name:
                continue
            if not path:
                return False, f"File '{name}' has no path."
            register_file_path(path, display_name=name)
            registered += 1
        return True, f"Registered {registered} file(s)"

    # input_source step (must run in main thread — modifies session_state)
    if step_type == "input_source":
        file_classes = cfg.get("file_classes", []) or []
        if not file_classes:
            return False, "No file classes defined"
        return run_input_source_step(file_classes)

    # Multi-file steps (import, append, header_map)
    if step_type in ("import", "append", "header_map"):
        src_label = cfg.get("src_file")
        src_path, err = _resolve(src_label, "source file")
        if err:
            return False, err
        tgt_path, err = _resolve(cfg.get("tgt_file"), "target file")
        if err:
            return False, err

        if step_type == "append":
            return run_append_step(src_path, tgt_path, cfg.get("src_sheet"), cfg.get("tgt_sheet"))

        if step_type == "import":
            groups = st.session_state.get("uploaded_file_groups") or {}
            all_src = groups.get(src_label) if src_label in groups else None
            if not all_src:
                all_src = [src_path]

            user_append = bool(cfg.get("append_mode", False))
            errors = []
            last_msg = ""
            for idx, sp in enumerate(all_src):
                mode = user_append if idx == 0 else True
                fn_col = cfg.get("filename_col") or None
                ok, msg = run_import_step(
                    sp, tgt_path, cfg.get("src_sheet"), cfg.get("tgt_sheet"),
                    cfg.get("mapping"),
                    src_header_row=int(cfg.get("src_header_row", 1) or 1),
                    tgt_header_row=int(cfg.get("tgt_header_row", 1) or 1),
                    append_mode=mode,
                    filters=cfg.get("filters") or [],
                    filter_combine=cfg.get("filter_combine", "AND"),
                    filename_col=fn_col,
                    filename_value=os.path.basename(sp) if fn_col else None,
                )
                if ok:
                    last_msg = msg
                else:
                    errors.append(f"{os.path.basename(sp)}: {msg}")

            if errors:
                partial = len(all_src) - len(errors)
                return partial > 0, f"Import: {partial}/{len(all_src)} files OK. Errors: {'; '.join(errors)}"
            suffix = f" ({len(all_src)} files)" if len(all_src) > 1 else ""
            return True, f"{last_msg}{suffix}"

        return run_header_mapping_step(
            src_path, tgt_path, cfg.get("src_sheet"), cfg.get("tgt_sheet"), cfg.get("mapping")
        )

    # Explicit VLOOKUP (4 file inputs)
    if step_type == "vlookup" and (cfg.get("mode") == "explicit" or cfg.get("lookup_value_file")):
        lv_path, err = _resolve(cfg.get("lookup_value_file"), "lookup value file")
        if err: return False, err
        s_path, err = _resolve(cfg.get("search_file"), "search file")
        if err: return False, err
        r_path, err = _resolve(cfg.get("return_file"), "return file")
        if err: return False, err
        d_path, err = _resolve(cfg.get("return_to_file"), "destination file")
        if err: return False, err

        fb_lookups = []
        for fb in cfg.get("fallback_lookups", []) or []:
            fb = fb or {}
            fb_s_path, err = _resolve(fb.get("search_file"), "fallback search file")
            if err: return False, err
            fb_r_path, err = _resolve(fb.get("return_file"), "fallback return file")
            if err: return False, err
            fb_lookups.append({
                "search_file_path": fb_s_path,
                "search_sheet_name": fb.get("search_sheet"),
                "search_header_row": int(fb.get("search_header_row", 1) or 1),
                "search_columns": fb.get("search_columns") or fb.get("search_column"),
                "return_file_path": fb_r_path,
                "return_sheet_name": fb.get("return_sheet"),
                "return_header_row": int(fb.get("return_header_row", 1) or 1),
                "return_column": fb.get("return_column"),
            })

        return run_vlookup_explicit_step(
            lookup_value_file_path=lv_path,
            lookup_value_sheet_name=cfg.get("lookup_value_sheet"),
            lookup_value_header_row=int(cfg.get("lookup_value_header_row", 1) or 1),
            lookup_value_column=cfg.get("lookup_value_column"),
            search_file_path=s_path,
            search_sheet_name=cfg.get("search_sheet"),
            search_header_row=int(cfg.get("search_header_row", 1) or 1),
            search_column=cfg.get("search_column"),
            return_file_path=r_path,
            return_sheet_name=cfg.get("return_sheet"),
            return_header_row=int(cfg.get("return_header_row", 1) or 1),
            return_column=cfg.get("return_column"),
            return_to_file_path=d_path,
            return_to_sheet_name=cfg.get("return_to_sheet"),
            return_to_header_row=int(cfg.get("return_to_header_row", 1) or 1),
            return_to_column=cfg.get("return_to_column"),
            not_found_value=cfg.get("not_found_value", ""),
            end_row=cfg.get("end_row", "last"),
            fallback_lookups=fb_lookups,
        )

    # Single-file operations
    file_label = cfg.get("file")
    file_path, err = _resolve(file_label, "file")
    if err:
        return False, err

    if step_type == "formula":
        return run_formula_step(
            file_path, cfg.get("sheet"), cfg.get("formula"),
            _resolve_column(file_path, cfg.get("sheet"), cfg),
            cfg.get("start"), cfg.get("end"),
            cfg.get("convert", False),
        )
    if step_type == "formula_broadcast":
        # all_sheets: uses the already-resolved file_path
        # all_files:  needs file group — executor uses single file_path only
        scope = cfg.get("scope", "all_sheets")
        incl = [s.strip() for s in (cfg.get("include_sheets") or "").split(",") if s.strip()]
        excl = [s.strip() for s in (cfg.get("exclude_sheets") or "").split(",") if s.strip()]
        return run_formula_broadcast_step(
            scope=scope,
            formula=cfg.get("formula", ""),
            cell=cfg.get("cell", ""),
            file_path=file_path,
            file_paths=[],          # single-file fallback; app layer passes full group
            sheet_name=cfg.get("sheet_name", ""),
            include_sheets=incl,
            exclude_sheets=excl,
        )
    if step_type == "copy_paste":
        return run_copy_paste_step(
            file_path, cfg.get("sheet"), cfg.get("src_col"),
            cfg.get("tgt_cols") if cfg.get("tgt_cols") else cfg.get("tgt_col"),
            cfg.get("paste_type"), cfg.get("header_row", 1),
        )
    if step_type == "convert_values":
        return run_convert_to_values_step(
            file_path, cfg.get("sheet"), cfg.get("column"),
            cfg.get("start"), cfg.get("end"),
        )
    if step_type == "forward_fill":
        return run_forward_fill_step(file_path, cfg.get("sheet"), cfg.get("column"))
    if step_type == "replace":
        return run_replace_single(
            file_path, cfg.get("sheet"), cfg.get("old_val"),
            cfg.get("new_val"), cfg.get("column"),
        )
    if step_type == "conditional":
        return run_conditional_write_step(
            file_path,
            cfg.get("sheet"),
            condition_col=cfg.get("cond_col"),
            condition_value=cfg.get("cond_val", ""),
            target_col=cfg.get("target_col"),
            write_value=cfg.get("write_val", ""),
            operator=cfg.get("operator", "equals"),
            header_row=int(cfg.get("header_row", 1) or 1),
            col_mode=cfg.get("col_mode", "letter"),
            rules=cfg.get("rules") or None,
            scope=cfg.get("scope", "single"),
            include_sheets=cfg.get("include_sheets") or [],
            exclude_sheets=cfg.get("exclude_sheets") or [],
        )
    if step_type == "write_cell":
        return run_write_cell_step(
            file_path, cfg.get("sheet"), cfg.get("mode", "single_cell"),
            cfg.get("value", ""),
            cell_ref=cfg.get("cell_ref", ""), column=cfg.get("column", ""),
            start_row=int(cfg.get("start_row", 2)),
            end_row=cfg.get("end_row", "last"),
            header_row=int(cfg.get("header_row", 1)),
        )
    if step_type == "filter":
        return run_filter_step(
            file_path, cfg.get("sheet"), cfg.get("column"), cfg.get("value"),
            remove_empty=bool(cfg.get("remove_empty", False)),
            header_row=int(cfg.get("header_row", 1) or 1),
        )
    if step_type == "sort":
        if isinstance(cfg.get("sort_cols"), list) and cfg.get("sort_cols"):
            cols = cfg.get("sort_cols") or []
            ascending = cfg.get("sort_ascending", True)
            header_row = int(cfg.get("header_row", 1) or 1)
        else:
            cols = [c.strip() for c in (cfg.get("columns", "") or "").split(",") if c.strip()]
            ascending = cfg.get("ascending", True)
            header_row = int(cfg.get("header_row", 1) or 1)
        return run_sort_step(file_path, cfg.get("sheet"), cols, ascending=ascending, header_row=header_row)
    if step_type == "insert":
        return run_insert_delete_step(
            file_path, cfg.get("sheet"), "insert",
            cfg.get("axis"), cfg.get("start"), cfg.get("end"),
        )
    if step_type == "delete":
        if cfg.get("row_mode") == "advanced_condition":
            return run_advanced_delete_step(
                file_path=file_path,
                sheet_scope=cfg.get("sheet_scope", "single"),
                sheet_name=cfg.get("sheet", ""),
                selected_sheets=cfg.get("selected_sheets") or [],
                include_sheets=cfg.get("include_sheets") or [],
                exclude_sheets=cfg.get("exclude_sheets") or [],
                filters=cfg.get("filters") or [],
                filter_combine=cfg.get("filter_combine", "AND"),
                header_row=int(cfg.get("header_row", 1) or 1),
                save_log=bool(cfg.get("save_log", True)),
                log_path=cfg.get("log_path", ""),
            )
        if cfg.get("row_mode") == "condition":
            return run_delete_by_condition_step(
                file_path, cfg.get("sheet"),
                filters=cfg.get("filters") or [],
                filter_combine=cfg.get("filter_combine", "AND"),
                header_row=int(cfg.get("header_row", 1) or 1),
            )
        if str(cfg.get("axis", "")).lower() == "column" and cfg.get("columns"):
            if cfg.get("action") == "clear_data":
                return run_clear_columns_data_step(
                    file_path, cfg.get("sheet"), cfg.get("columns"),
                    header_row=int(cfg.get("header_row", 1) or 1),
                    end_row=cfg.get("end_row", "last"),
                )
            return run_delete_columns_step(
                file_path, cfg.get("sheet"), cfg.get("columns"),
                header_row=int(cfg.get("header_row", 1) or 1),
            )
        return run_insert_delete_step(
            file_path, cfg.get("sheet"), "delete",
            cfg.get("axis"), cfg.get("start"), cfg.get("end"),
        )
    if step_type == "merge_cells":
        return run_merge_cells_step(file_path, cfg.get("sheet"), cfg.get("range"))
    if step_type == "delete_sheets":
        return run_delete_sheets_step(
            file_path,
            mode=cfg.get("mode", "named"),
            sheet_names=cfg.get("sheet_names") or [],
            pattern=cfg.get("pattern", ""),
            pattern_type=cfg.get("pattern_type", "contains"),
            case_sensitive=bool(cfg.get("case_sensitive", False)),
            protect_last=bool(cfg.get("protect_last", True)),
        )
    if step_type == "format":
        return run_format_step(
            file_path,
            cfg.get("sheet"),
            format_mode=cfg.get("format_mode", "static"),
            header_row=int(cfg.get("header_row", 1) or 1),
            col_mode=cfg.get("col_mode", "letter"),
            static_rules=cfg.get("static_rules") or [],
            cond_rules=cfg.get("cond_rules") or [],
        )
    if step_type == "sumifs":
        src_path_si, _err = _resolve(cfg.get("src_workbook"), "src_workbook")
        if _err: src_path_si = cfg.get("src_workbook") or ""
        return run_sumifs_step(
            target_file_path=file_path,
            target_sheet=cfg.get("sheet"),
            source_file_path=src_path_si,
            source_sheet=cfg.get("src_sheet"),
            src_lookup_col=cfg.get("src_lookup_col") or cfg.get("sum_range", "A"),
            tgt_lookup_col=cfg.get("tgt_lookup_col", "A"),
            sum_col=cfg.get("sum_col") or cfg.get("sum_range", "B"),
            output_col=cfg.get("output_col", "B"),
            src_header_row=int(cfg.get("src_header_row", 1) or 1),
            tgt_header_row=int(cfg.get("tgt_header_row", 1) or 1),
        )
    if step_type == "advance_vlookup":
        if cfg.get("lookup_value_columns"):
            from operations import run_advance_vlookup_v2
            v2_config = dict(cfg)
            for fld in ("condition_file", "lookup_value_file", "search_file", "return_file", "destination_file"):
                v2_config[fld], err = _resolve(cfg.get(fld), fld)
                if err: return False, err
            if cfg.get("else_source_file"):
                v2_config["else_source_file"], err = _resolve(cfg.get("else_source_file"), "else_source_file")
                if err: return False, err
            return run_advance_vlookup_v2(v2_config)
        else:
            lookup_path = get_file_path(cfg.get("lookup_file", ""))
            return run_advance_vlookup_step(
                file_path, cfg.get("sheet"), cfg.get("prefix"),
                cfg.get("src_col"), lookup_path, cfg.get("lookup_sheet"),
                cfg.get("key_col"), cfg.get("return_col"), cfg.get("output_col"),
                fallback_value=cfg.get("fallback_value"),
                fallback_col=cfg.get("fallback_col"),
                target_header_row=int(cfg.get("target_header_row", 1) or 1),
                lookup_header_row=int(cfg.get("lookup_header_row", 1) or 1),
                condition_enabled=bool(cfg.get("condition_enabled", False)),
                condition_column=cfg.get("condition_column"),
                condition_value=cfg.get("condition_value"),
                else_source_column=cfg.get("else_source_column"),
            )
    if step_type == "vlookup":
        lookup_path = get_file_path(cfg.get("lookup_file", ""))
        return run_vlookup_step(
            file_path, cfg.get("sheet"), cfg.get("main_col"),
            [{"file_path": lookup_path, "sheet_name": cfg.get("lookup_sheet"),
              "key_col": cfg.get("key_col"), "return_col": cfg.get("return_col")}],
            cfg.get("fallback"),
        )
    if step_type == "data_input":
        success, msg, _df = run_data_input_step(file_path, cfg.get("sheet"))
        return success, msg

    if step_type == "normalize_import":
        src_label = cfg.get("src_file")
        src_path, err = _resolve(src_label, "source file")
        if err:
            return False, err
        tgt_path, err = _resolve(cfg.get("tgt_file"), "target file")
        if err:
            return False, err
        # Use the full file group when the class came from a folder scan
        groups = st.session_state.get("uploaded_file_groups") or {}
        src_paths = groups.get(src_label) if src_label in groups else None
        return run_normalize_import_step(
            src_path=src_path,
            src_sheet=cfg.get("src_sheet"),
            src_header_row=int(cfg.get("src_header_row", 1) or 1),
            mapping=cfg.get("mapping", []) or [],
            tgt_path=tgt_path,
            tgt_sheet=cfg.get("tgt_sheet"),
            tgt_header_row=int(cfg.get("tgt_header_row", 1) or 1),
            append_mode=bool(cfg.get("append_mode", True)),
            src_paths=src_paths,
        )

    if step_type == "pivot":
        src_path, err = _resolve(cfg.get("src_file"), "src_file")
        if err: return False, err
        tgt_path, err = _resolve(cfg.get("tgt_file"), "tgt_file")
        if err: return False, err
        return run_pivot_step(
            src_file_path=src_path,
            src_sheet=cfg.get("src_sheet"),
            src_header_row=int(cfg.get("src_header_row", 1) or 1),
            tgt_file_path=tgt_path,
            tgt_sheet=cfg.get("tgt_sheet", "Pivot"),
            tgt_start_row=int(cfg.get("tgt_start_row", 1) or 1),
            tgt_start_col=cfg.get("tgt_start_col", "A"),
            overwrite_sheet=bool(cfg.get("overwrite_sheet", True)),
            row_fields=cfg.get("row_fields") or [],
            col_fields=cfg.get("col_fields") or [],
            value_fields=cfg.get("value_fields") or [],
            filter_conditions=cfg.get("filter_conditions") or [],
            filter_combine=cfg.get("filter_combine", "AND"),
            fill_value=cfg.get("fill_value", 0),
            grand_total_rows=bool(cfg.get("grand_total_rows", False)),
            grand_total_cols=bool(cfg.get("grand_total_cols", False)),
            sort_field=cfg.get("sort_field", ""),
            sort_ascending=bool(cfg.get("sort_ascending", True)),
            apply_formatting=bool(cfg.get("apply_formatting", True)),
            header_bg_color=cfg.get("header_bg_color", "1F3864"),
            header_font_color=cfg.get("header_font_color", "FFFFFF"),
            total_bg_color=cfg.get("total_bg_color", "D9E1F2"),
        )

    if step_type == "folder_summary":
        src_label = cfg.get("src_file")
        src_path, err = _resolve(src_label, "source file")
        if err:
            return False, err
        tgt_path, err2 = _resolve(cfg.get("tgt_file"), "target file")
        if err2:
            return False, err2
        # Resolve full file list from group if available
        groups = st.session_state.get("uploaded_file_groups") or {}
        src_paths = groups.get(src_label) if src_label in (groups or {}) else None
        if not src_paths:
            src_paths = [src_path]
        return run_folder_summary_step(
            src_file_paths=src_paths,
            src_sheet=cfg.get("src_sheet", ""),
            header_row=int(cfg.get("header_row", 1) or 1),
            source_mode=cfg.get("source_mode", "folder"),
            sheet_selection=cfg.get("sheet_selection", "all"),
            sheet_pattern=cfg.get("sheet_pattern", ""),
            sheet_list=cfg.get("sheet_list") or [],
            sheet_exclude=cfg.get("sheet_exclude", ""),
            variables=cfg.get("variables") or [],
            formulas=cfg.get("formulas") or [],
            tgt_file=tgt_path,
            tgt_sheet=cfg.get("tgt_sheet", "Summary"),
            tgt_header_row=int(cfg.get("tgt_header_row", 1) or 1),
            filename_col=cfg.get("filename_col", "File Name"),
            sheet_name_col=cfg.get("sheet_name_col", ""),
            append_mode=bool(cfg.get("append_mode", True)),
            include_var_cols=bool(cfg.get("include_var_cols", True)),
            include_status_col=bool(cfg.get("include_status_col", True)),
            status_col_name=cfg.get("status_col_name", "Status"),
            on_missing_default=cfg.get("on_missing_default", 0),
        )

    if step_type == "split_export":
        src_path, err = _resolve(cfg.get("src_file"), "src_file")
        if err: return False, err
        tgt_label = cfg.get("tgt_file")
        tgt_path = get_file_path(tgt_label) if tgt_label else ""
        return run_split_export_step(
            src_file_path=src_path,
            src_sheet=cfg.get("src_sheet", ""),
            src_header_row=int(cfg.get("src_header_row", 1) or 1),
            split_mode=cfg.get("split_mode", "single"),
            split_col=cfg.get("split_col", ""),
            split_col2=cfg.get("split_col2", ""),
            output_folder=cfg.get("output_folder", ""),
            tgt_file=tgt_path or cfg.get("tgt_file_path", ""),
            tgt_sheet=cfg.get("tgt_sheet", "Export"),
            file_name_template=cfg.get("file_name_template", "{value}.xlsx"),
            sheet_name_template=cfg.get("sheet_name_template", "{value}"),
            sheet_name_template2=cfg.get("sheet_name_template2", "{value}"),
            filter_conditions=cfg.get("filter_conditions") or [],
            filter_combine=cfg.get("filter_combine", "AND"),
            include_columns=cfg.get("include_columns") or [],
            exclude_columns=cfg.get("exclude_columns") or [],
            sort_cols=cfg.get("sort_cols") or [],
            sort_ascending=cfg.get("sort_ascending", True),
            include_header=bool(cfg.get("include_header", True)),
            overwrite=bool(cfg.get("overwrite", True)),
            max_rows_per_sheet=int(cfg.get("max_rows_per_sheet", 0) or 0),
            add_summary_sheet=bool(cfg.get("add_summary_sheet", False)),
            summary_sheet_name=cfg.get("summary_sheet_name", "_Summary"),
            apply_formatting=bool(cfg.get("apply_formatting", True)),
            header_bg_color=cfg.get("header_bg_color", "1F3864"),
            header_font_color=cfg.get("header_font_color", "FFFFFF"),
        )

    if step_type == "unpivot":
        src_path, err = _resolve(cfg.get("src_file"), "src_file")
        if err: return False, err
        tgt_label = cfg.get("tgt_file")
        tgt_path = get_file_path(tgt_label) if tgt_label else ""
        return run_unpivot_step(
            src_file_path=src_path,
            src_sheet=cfg.get("src_sheet", ""),
            src_header_row=int(cfg.get("src_header_row", 1) or 1),
            id_cols=cfg.get("id_cols") or [],
            value_cols=cfg.get("value_cols") or [],
            var_name=cfg.get("var_name", "Category"),
            value_name=cfg.get("value_name", "Value"),
            tgt_file=tgt_path or "",
            tgt_sheet=cfg.get("tgt_sheet", "Unpivot"),
            append_mode=bool(cfg.get("append_mode", False)),
            drop_na=bool(cfg.get("drop_na", True)),
        )

    if step_type == "sheet_updater":
        src_path, err = _resolve(cfg.get("src_file"), "src_file")
        if err: return False, err
        tgt_label = cfg.get("tgt_file")
        tgt_path = get_file_path(tgt_label) if tgt_label else ""
        return run_sheet_updater_step(
            src_file_path=src_path,
            src_sheet=cfg.get("src_sheet", ""),
            src_col_category=cfg.get("src_col_category", ""),
            src_col_particulars=cfg.get("src_col_particulars", ""),
            src_col_value=cfg.get("src_col_value", ""),
            tgt_file=tgt_path or "",
            sheet_match_mode=cfg.get("sheet_match_mode", "cat_in_sheet"),
            tgt_col_lookup=cfg.get("tgt_col_lookup", ""),
            tgt_col_write=cfg.get("tgt_col_write", ""),
            tgt_header_row=int(cfg.get("tgt_header_row", 1) or 1),
            case_sensitive=bool(cfg.get("case_sensitive", False)),
            particulars_match_mode=cfg.get("particulars_match_mode", "iexact"),
            particulars_rules=cfg.get("particulars_rules") or [],
        )

    return False, f"Unknown operation: {step_type}"


def execute_step_with_file_map(step, file_map):
    """
    Thread-safe step execution without Streamlit APIs.
    file_map is a {label -> absolute_path} snapshot.
    """
    step_type = (step.get("type") or "").lower().strip()
    cfg = step.get("config", {}) or {}

    def path_for(label):
        if not label:
            return ""
        return str(file_map.get(label, "") or "")

    def resolve(label, field):
        if not label:
            return "", f"Missing file for '{field}'."
        p = path_for(label)
        if p:
            return p, None
        labels = sorted([str(k) for k in file_map.keys()])
        preview = ", ".join([f"'{x}'" for x in labels[:10]])
        return "", f"File not found: '{label}' for '{field}'. Registered: [{preview}]"

    if step_type in ("add_files", "input_source"):
        return False, f"{step_type} must run in main thread"

    # Multi-file steps
    if step_type in ("import", "append", "header_map"):
        src_path, err = resolve(cfg.get("src_file"), "source file")
        if err: return False, err
        tgt_path, err = resolve(cfg.get("tgt_file"), "target file")
        if err: return False, err

        if step_type == "append":
            return run_append_step(src_path, tgt_path, cfg.get("src_sheet"), cfg.get("tgt_sheet"))
        if step_type == "import":
            fn_col = cfg.get("filename_col") or None
            return run_import_step(
                src_path, tgt_path, cfg.get("src_sheet"), cfg.get("tgt_sheet"),
                cfg.get("mapping"),
                src_header_row=int(cfg.get("src_header_row", 1) or 1),
                tgt_header_row=int(cfg.get("tgt_header_row", 1) or 1),
                append_mode=bool(cfg.get("append_mode", False)),
                filters=cfg.get("filters") or [],
                filter_combine=cfg.get("filter_combine", "AND"),
                filename_col=fn_col,
                filename_value=os.path.basename(src_path) if fn_col else None,
            )
        return run_header_mapping_step(
            src_path, tgt_path, cfg.get("src_sheet"), cfg.get("tgt_sheet"), cfg.get("mapping")
        )

    # Explicit VLOOKUP
    if step_type == "vlookup" and (cfg.get("mode") == "explicit" or cfg.get("lookup_value_file")):
        lv_path, err = resolve(cfg.get("lookup_value_file"), "lookup value file")
        if err: return False, err
        s_path, err = resolve(cfg.get("search_file"), "search file")
        if err: return False, err
        r_path, err = resolve(cfg.get("return_file"), "return file")
        if err: return False, err
        d_path, err = resolve(cfg.get("return_to_file"), "destination file")
        if err: return False, err

        fb_lookups = []
        for fb in cfg.get("fallback_lookups", []) or []:
            fb = fb or {}
            fb_s_path, err = resolve(fb.get("search_file"), "fallback search file")
            if err: return False, err
            fb_r_path, err = resolve(fb.get("return_file"), "fallback return file")
            if err: return False, err
            fb_lookups.append({
                "search_file_path": fb_s_path, "search_sheet_name": fb.get("search_sheet"),
                "search_header_row": int(fb.get("search_header_row", 1) or 1),
                "search_columns": fb.get("search_columns") or fb.get("search_column"),
                "return_file_path": fb_r_path, "return_sheet_name": fb.get("return_sheet"),
                "return_header_row": int(fb.get("return_header_row", 1) or 1),
                "return_column": fb.get("return_column"),
            })

        return run_vlookup_explicit_step(
            lookup_value_file_path=lv_path,
            lookup_value_sheet_name=cfg.get("lookup_value_sheet"),
            lookup_value_header_row=int(cfg.get("lookup_value_header_row", 1) or 1),
            lookup_value_column=cfg.get("lookup_value_column"),
            search_file_path=s_path, search_sheet_name=cfg.get("search_sheet"),
            search_header_row=int(cfg.get("search_header_row", 1) or 1),
            search_column=cfg.get("search_column"),
            return_file_path=r_path, return_sheet_name=cfg.get("return_sheet"),
            return_header_row=int(cfg.get("return_header_row", 1) or 1),
            return_column=cfg.get("return_column"),
            return_to_file_path=d_path, return_to_sheet_name=cfg.get("return_to_sheet"),
            return_to_header_row=int(cfg.get("return_to_header_row", 1) or 1),
            return_to_column=cfg.get("return_to_column"),
            not_found_value=cfg.get("not_found_value", ""),
            end_row=cfg.get("end_row", "last"),
            fallback_lookups=fb_lookups,
        )

    # Single-file operations
    file_label = cfg.get("file")
    file_path, err = resolve(file_label, "file")
    if err: return False, err

    if step_type == "formula":
        return run_formula_step(file_path, cfg.get("sheet"), cfg.get("formula"), _resolve_column(file_path, cfg.get("sheet"), cfg), cfg.get("start"), cfg.get("end"), cfg.get("convert", False))
    if step_type == "formula_broadcast":
        scope = cfg.get("scope", "all_sheets")
        incl = [s.strip() for s in (cfg.get("include_sheets") or "").split(",") if s.strip()]
        excl = [s.strip() for s in (cfg.get("exclude_sheets") or "").split(",") if s.strip()]
        fp_group = [path_for(lbl) for lbl in (cfg.get("file_paths_labels") or []) if path_for(lbl)]
        return run_formula_broadcast_step(
            scope=scope,
            formula=cfg.get("formula", ""),
            cell=cfg.get("cell", ""),
            file_path=file_path,
            file_paths=fp_group,
            sheet_name=cfg.get("sheet_name", ""),
            include_sheets=incl,
            exclude_sheets=excl,
        )
    if step_type == "copy_paste":
        return run_copy_paste_step(file_path, cfg.get("sheet"), cfg.get("src_col"), cfg.get("tgt_cols") if cfg.get("tgt_cols") else cfg.get("tgt_col"), cfg.get("paste_type"), cfg.get("header_row", 1))
    if step_type == "convert_values":
        return run_convert_to_values_step(file_path, cfg.get("sheet"), cfg.get("column"), cfg.get("start"), cfg.get("end"))
    if step_type == "forward_fill":
        return run_forward_fill_step(file_path, cfg.get("sheet"), cfg.get("column"))
    if step_type == "replace":
        return run_replace_single(file_path, cfg.get("sheet"), cfg.get("old_val"), cfg.get("new_val"), cfg.get("column"))
    if step_type == "conditional":
        return run_conditional_write_step(
            file_path,
            cfg.get("sheet"),
            condition_col=cfg.get("cond_col"),
            condition_value=cfg.get("cond_val", ""),
            target_col=cfg.get("target_col"),
            write_value=cfg.get("write_val", ""),
            operator=cfg.get("operator", "equals"),
            header_row=int(cfg.get("header_row", 1) or 1),
            col_mode=cfg.get("col_mode", "letter"),
            rules=cfg.get("rules") or None,
            scope=cfg.get("scope", "single"),
            include_sheets=cfg.get("include_sheets") or [],
            exclude_sheets=cfg.get("exclude_sheets") or [],
        )
    if step_type == "write_cell":
        return run_write_cell_step(file_path, cfg.get("sheet"), cfg.get("mode", "single_cell"), cfg.get("value", ""), cell_ref=cfg.get("cell_ref", ""), column=cfg.get("column", ""), start_row=int(cfg.get("start_row", 2)), end_row=cfg.get("end_row", "last"), header_row=int(cfg.get("header_row", 1)))
    if step_type == "filter":
        return run_filter_step(file_path, cfg.get("sheet"), cfg.get("column"), cfg.get("value"), remove_empty=bool(cfg.get("remove_empty", False)), header_row=int(cfg.get("header_row", 1) or 1))
    if step_type == "sort":
        if isinstance(cfg.get("sort_cols"), list) and cfg.get("sort_cols"):
            cols, ascending, header_row = cfg["sort_cols"], cfg.get("sort_ascending", True), int(cfg.get("header_row", 1) or 1)
        else:
            cols = [c.strip() for c in (cfg.get("columns", "") or "").split(",") if c.strip()]
            ascending, header_row = cfg.get("ascending", True), int(cfg.get("header_row", 1) or 1)
        return run_sort_step(file_path, cfg.get("sheet"), cols, ascending=ascending, header_row=header_row)
    if step_type == "insert":
        return run_insert_delete_step(file_path, cfg.get("sheet"), "insert", cfg.get("axis"), cfg.get("start"), cfg.get("end"))
    if step_type == "delete":
        if cfg.get("row_mode") == "advanced_condition":
            return run_advanced_delete_step(
                file_path=file_path,
                sheet_scope=cfg.get("sheet_scope", "single"),
                sheet_name=cfg.get("sheet", ""),
                selected_sheets=cfg.get("selected_sheets") or [],
                include_sheets=cfg.get("include_sheets") or [],
                exclude_sheets=cfg.get("exclude_sheets") or [],
                filters=cfg.get("filters") or [],
                filter_combine=cfg.get("filter_combine", "AND"),
                header_row=int(cfg.get("header_row", 1) or 1),
                save_log=bool(cfg.get("save_log", True)),
                log_path=cfg.get("log_path", ""),
            )
        if cfg.get("row_mode") == "condition":
            return run_delete_by_condition_step(file_path, cfg.get("sheet"), filters=cfg.get("filters") or [], filter_combine=cfg.get("filter_combine", "AND"), header_row=int(cfg.get("header_row", 1) or 1))
        if str(cfg.get("axis", "")).lower() == "column" and cfg.get("columns"):
            if cfg.get("action") == "clear_data":
                return run_clear_columns_data_step(file_path, cfg.get("sheet"), cfg.get("columns"), header_row=int(cfg.get("header_row", 1) or 1), end_row=cfg.get("end_row", "last"))
            return run_delete_columns_step(file_path, cfg.get("sheet"), cfg.get("columns"), header_row=int(cfg.get("header_row", 1) or 1))
        return run_insert_delete_step(file_path, cfg.get("sheet"), "delete", cfg.get("axis"), cfg.get("start"), cfg.get("end"))
    if step_type == "merge_cells":
        return run_merge_cells_step(file_path, cfg.get("sheet"), cfg.get("range"))
    if step_type == "delete_sheets":
        return run_delete_sheets_step(
            file_path,
            mode=cfg.get("mode", "named"),
            sheet_names=cfg.get("sheet_names") or [],
            pattern=cfg.get("pattern", ""),
            pattern_type=cfg.get("pattern_type", "contains"),
            case_sensitive=bool(cfg.get("case_sensitive", False)),
            protect_last=bool(cfg.get("protect_last", True)),
        )
    if step_type == "format":
        return run_format_step(
            file_path,
            cfg.get("sheet"),
            format_mode=cfg.get("format_mode", "static"),
            header_row=int(cfg.get("header_row", 1) or 1),
            col_mode=cfg.get("col_mode", "letter"),
            static_rules=cfg.get("static_rules") or [],
            cond_rules=cfg.get("cond_rules") or [],
        )
    if step_type == "sumifs":
        src_path_si = resolve(cfg.get("src_workbook"), "src_workbook")[0] or cfg.get("src_workbook") or ""
        return run_sumifs_step(
            target_file_path=file_path,
            target_sheet=cfg.get("sheet"),
            source_file_path=src_path_si,
            source_sheet=cfg.get("src_sheet"),
            src_lookup_col=cfg.get("src_lookup_col") or cfg.get("sum_range", "A"),
            tgt_lookup_col=cfg.get("tgt_lookup_col", "A"),
            sum_col=cfg.get("sum_col") or cfg.get("sum_range", "B"),
            output_col=cfg.get("output_col", "B"),
            src_header_row=int(cfg.get("src_header_row", 1) or 1),
            tgt_header_row=int(cfg.get("tgt_header_row", 1) or 1),
        )
    if step_type == "advance_vlookup":
        if cfg.get("lookup_value_columns"):
            from operations import run_advance_vlookup_v2
            v2_config = dict(cfg)
            for fld in ("condition_file", "lookup_value_file", "search_file", "return_file", "destination_file"):
                v2_config[fld], err = resolve(cfg.get(fld), fld)
                if err: return False, err
            if cfg.get("else_source_file"):
                v2_config["else_source_file"], err = resolve(cfg.get("else_source_file"), "else_source_file")
                if err: return False, err
            return run_advance_vlookup_v2(v2_config)
        else:
            lookup_path = path_for(cfg.get("lookup_file"))
            return run_advance_vlookup_step(file_path, cfg.get("sheet"), cfg.get("prefix"), cfg.get("src_col"), lookup_path, cfg.get("lookup_sheet"), cfg.get("key_col"), cfg.get("return_col"), cfg.get("output_col"), fallback_value=cfg.get("fallback_value"), fallback_col=cfg.get("fallback_col"), target_header_row=int(cfg.get("target_header_row", 1) or 1), lookup_header_row=int(cfg.get("lookup_header_row", 1) or 1), condition_enabled=bool(cfg.get("condition_enabled", False)), condition_column=cfg.get("condition_column"), condition_value=cfg.get("condition_value"), else_source_column=cfg.get("else_source_column"))
    if step_type == "vlookup":
        lookup_path = path_for(cfg.get("lookup_file"))
        return run_vlookup_step(file_path, cfg.get("sheet"), cfg.get("main_col"), [{"file_path": lookup_path, "sheet_name": cfg.get("lookup_sheet"), "key_col": cfg.get("key_col"), "return_col": cfg.get("return_col")}], cfg.get("fallback"))
    if step_type == "data_input":
        success, msg, _df = run_data_input_step(file_path, cfg.get("sheet"))
        return success, msg
    if step_type == "normalize_import":
        src_label = cfg.get("src_file")
        src_path = path_for(src_label)
        tgt_path = path_for(cfg.get("tgt_file"))
        if not src_path:
            return False, f"Source file not found: {src_label}"
        if not tgt_path:
            return False, f"Target file not found: {cfg.get('tgt_file')}"
        # file_map snapshot doesn't carry groups — pass None; single-file fallback applies
        src_paths = None
        return run_normalize_import_step(
            src_path=src_path,
            src_sheet=cfg.get("src_sheet"),
            src_header_row=int(cfg.get("src_header_row", 1) or 1),
            mapping=cfg.get("mapping", []) or [],
            tgt_path=tgt_path,
            tgt_sheet=cfg.get("tgt_sheet"),
            tgt_header_row=int(cfg.get("tgt_header_row", 1) or 1),
            append_mode=bool(cfg.get("append_mode", True)),
            src_paths=src_paths,
        )

    if step_type == "pivot":
        src_p = path_for(cfg.get("src_file"))
        tgt_p = path_for(cfg.get("tgt_file"))
        if not src_p: return False, f"Source file not found: {cfg.get('src_file')}"
        if not tgt_p: return False, f"Target file not found: {cfg.get('tgt_file')}"
        return run_pivot_step(
            src_file_path=src_p,
            src_sheet=cfg.get("src_sheet"),
            src_header_row=int(cfg.get("src_header_row", 1) or 1),
            tgt_file_path=tgt_p,
            tgt_sheet=cfg.get("tgt_sheet", "Pivot"),
            tgt_start_row=int(cfg.get("tgt_start_row", 1) or 1),
            tgt_start_col=cfg.get("tgt_start_col", "A"),
            overwrite_sheet=bool(cfg.get("overwrite_sheet", True)),
            row_fields=cfg.get("row_fields") or [],
            col_fields=cfg.get("col_fields") or [],
            value_fields=cfg.get("value_fields") or [],
            filter_conditions=cfg.get("filter_conditions") or [],
            filter_combine=cfg.get("filter_combine", "AND"),
            fill_value=cfg.get("fill_value", 0),
            grand_total_rows=bool(cfg.get("grand_total_rows", False)),
            grand_total_cols=bool(cfg.get("grand_total_cols", False)),
            sort_field=cfg.get("sort_field", ""),
            sort_ascending=bool(cfg.get("sort_ascending", True)),
            apply_formatting=bool(cfg.get("apply_formatting", True)),
            header_bg_color=cfg.get("header_bg_color", "1F3864"),
            header_font_color=cfg.get("header_font_color", "FFFFFF"),
            total_bg_color=cfg.get("total_bg_color", "D9E1F2"),
        )

    if step_type == "folder_summary":
        src_label = cfg.get("src_file")
        src_path = path_for(src_label)
        tgt_path = path_for(cfg.get("tgt_file"))
        if not src_path:
            return False, f"Source file not found: {src_label}"
        if not tgt_path:
            return False, f"Target file not found: {cfg.get('tgt_file')}"
        src_paths = [src_path]
        return run_folder_summary_step(
            src_file_paths=src_paths,
            src_sheet=cfg.get("src_sheet", ""),
            header_row=int(cfg.get("header_row", 1) or 1),
            source_mode=cfg.get("source_mode", "folder"),
            sheet_selection=cfg.get("sheet_selection", "all"),
            sheet_pattern=cfg.get("sheet_pattern", ""),
            sheet_list=cfg.get("sheet_list") or [],
            sheet_exclude=cfg.get("sheet_exclude", ""),
            variables=cfg.get("variables") or [],
            formulas=cfg.get("formulas") or [],
            tgt_file=tgt_path,
            tgt_sheet=cfg.get("tgt_sheet", "Summary"),
            tgt_header_row=int(cfg.get("tgt_header_row", 1) or 1),
            filename_col=cfg.get("filename_col", "File Name"),
            sheet_name_col=cfg.get("sheet_name_col", ""),
            append_mode=bool(cfg.get("append_mode", True)),
            include_var_cols=bool(cfg.get("include_var_cols", True)),
            include_status_col=bool(cfg.get("include_status_col", True)),
            status_col_name=cfg.get("status_col_name", "Status"),
            on_missing_default=cfg.get("on_missing_default", 0),
        )

    if step_type == "unpivot":
        src_p = path_for(cfg.get("src_file"))
        if not src_p: return False, f"Source file not found: {cfg.get('src_file')}"
        tgt_p = path_for(cfg.get("tgt_file")) or ""
        return run_unpivot_step(
            src_file_path=src_p,
            src_sheet=cfg.get("src_sheet", ""),
            src_header_row=int(cfg.get("src_header_row", 1) or 1),
            id_cols=cfg.get("id_cols") or [],
            value_cols=cfg.get("value_cols") or [],
            var_name=cfg.get("var_name", "Category"),
            value_name=cfg.get("value_name", "Value"),
            tgt_file=tgt_p,
            tgt_sheet=cfg.get("tgt_sheet", "Unpivot"),
            append_mode=bool(cfg.get("append_mode", False)),
            drop_na=bool(cfg.get("drop_na", True)),
        )

    if step_type == "sheet_updater":
        src_p = path_for(cfg.get("src_file"))
        if not src_p: return False, f"Source file not found: {cfg.get('src_file')}"
        tgt_p = path_for(cfg.get("tgt_file")) or ""
        return run_sheet_updater_step(
            src_file_path=src_p,
            src_sheet=cfg.get("src_sheet", ""),
            src_col_category=cfg.get("src_col_category", ""),
            src_col_particulars=cfg.get("src_col_particulars", ""),
            src_col_value=cfg.get("src_col_value", ""),
            tgt_file=tgt_p,
            sheet_match_mode=cfg.get("sheet_match_mode", "cat_in_sheet"),
            tgt_col_lookup=cfg.get("tgt_col_lookup", ""),
            tgt_col_write=cfg.get("tgt_col_write", ""),
            tgt_header_row=int(cfg.get("tgt_header_row", 1) or 1),
            case_sensitive=bool(cfg.get("case_sensitive", False)),
            particulars_match_mode=cfg.get("particulars_match_mode", "iexact"),
            particulars_rules=cfg.get("particulars_rules") or [],
        )

    if step_type == "split_export":
        src_p = path_for(cfg.get("src_file"))
        if not src_p: return False, f"Source file not found: {cfg.get('src_file')}"
        tgt_label = cfg.get("tgt_file")
        tgt_p = path_for(tgt_label) if tgt_label else ""
        return run_split_export_step(
            src_file_path=src_p,
            src_sheet=cfg.get("src_sheet", ""),
            src_header_row=int(cfg.get("src_header_row", 1) or 1),
            split_mode=cfg.get("split_mode", "single"),
            split_col=cfg.get("split_col", ""),
            split_col2=cfg.get("split_col2", ""),
            output_folder=cfg.get("output_folder", ""),
            tgt_file=tgt_p or cfg.get("tgt_file_path", ""),
            tgt_sheet=cfg.get("tgt_sheet", "Export"),
            file_name_template=cfg.get("file_name_template", "{value}.xlsx"),
            sheet_name_template=cfg.get("sheet_name_template", "{value}"),
            sheet_name_template2=cfg.get("sheet_name_template2", "{value}"),
            filter_conditions=cfg.get("filter_conditions") or [],
            filter_combine=cfg.get("filter_combine", "AND"),
            include_columns=cfg.get("include_columns") or [],
            exclude_columns=cfg.get("exclude_columns") or [],
            sort_cols=cfg.get("sort_cols") or [],
            sort_ascending=cfg.get("sort_ascending", True),
            include_header=bool(cfg.get("include_header", True)),
            overwrite=bool(cfg.get("overwrite", True)),
            max_rows_per_sheet=int(cfg.get("max_rows_per_sheet", 0) or 0),
            add_summary_sheet=bool(cfg.get("add_summary_sheet", False)),
            summary_sheet_name=cfg.get("summary_sheet_name", "_Summary"),
            apply_formatting=bool(cfg.get("apply_formatting", True)),
            header_bg_color=cfg.get("header_bg_color", "1F3864"),
            header_font_color=cfg.get("header_font_color", "FFFFFF"),
        )

    return False, f"Unknown operation: {step_type}"


def execute_workflow_sequential(start_step=1, end_step=None):
    """Execute workflow steps sequentially with progress tracking."""
    from core.error_translator import translate_error

    progress = st.progress(0)
    status = st.empty()

    total = len(st.session_state.workflow_steps)
    if end_step is None:
        end_step = total

    start_idx = max(0, int(start_step) - 1)
    end_idx = min(total - 1, int(end_step) - 1)
    if start_idx > end_idx:
        st.warning("Invalid step range.")
        return []

    steps_to_run = list(enumerate(st.session_state.workflow_steps))[start_idx:end_idx + 1]
    range_total = len(steps_to_run)
    results = []

    wb_cache.activate()
    try:
        for j, (i, step) in enumerate(steps_to_run):
            from core.operation_registry import get_user_label
            step_type = step.get("type", "")
            label = get_user_label(step_type)
            status.text(f"Step {i + 1}/{total}: {label}")

            try:
                swapped_paths, original_mappings = prepare_non_destructive_export(i + 1, step)
                success, msg = execute_step(step)

                if swapped_paths:
                    finalize_non_destructive_export(i + 1, swapped_paths, original_mappings)

                results.append({
                    "step_num": i + 1,
                    "name": step.get("name", ""),
                    "type": step_type,
                    "success": success,
                    "message": msg,
                    "exported": bool(swapped_paths),
                })

                if success:
                    export_note = " (exported, original unchanged)" if swapped_paths else ""
                    st.success(f"Step {i + 1}: {msg}{export_note}")
                else:
                    translated = translate_error(msg, step_type)
                    st.error(f"Step {i + 1} - {translated['title']}")
                    st.markdown(f"**What happened:** {translated['message']}")
                    if translated.get("suggestion"):
                        st.info(f"**Suggestion:** {translated['suggestion']}")
                    with st.expander("Technical details"):
                        st.code(translated["technical_detail"])
                    break
            except Exception as e:
                translated = translate_error(str(e), step_type)
                st.error(f"Step {i + 1} - {translated['title']}")
                st.markdown(f"**What happened:** {translated['message']}")
                if translated.get("suggestion"):
                    st.info(f"**Suggestion:** {translated['suggestion']}")
                results.append({
                    "step_num": i + 1, "name": step.get("name", ""),
                    "type": step_type, "success": False, "message": str(e),
                })
                break

            progress.progress((j + 1) / range_total)

        status.text("Complete!")
    finally:
        wb_cache.deactivate()
    return results


def execute_workflow_dag(start_step=1, end_step=None):
    """Execute steps using dependency-aware parallel scheduling (DAG)."""
    from core.error_translator import translate_error
    from core.operation_registry import get_user_label

    ensure_step_ids()
    steps_all = st.session_state.get("workflow_steps", []) or []
    total = len(steps_all)
    if total == 0:
        st.warning("No steps to execute.")
        return []

    if end_step is None:
        end_step = total
    start_idx = max(0, int(start_step) - 1)
    end_idx = min(total - 1, int(end_step) - 1)
    if start_idx > end_idx:
        st.warning("Invalid step range.")
        return []

    selected_ids = {steps_all[i].get("id") for i in range(start_idx, end_idx + 1) if steps_all[i].get("id")}
    id_to_step = {s.get("id"): s for s in steps_all if s.get("id")}
    id_to_index = {s.get("id"): i for i, s in enumerate(steps_all) if s.get("id")}

    # Collect all needed nodes (selected + upstream deps)
    needed = set(selected_ids)
    stack = list(selected_ids)
    while stack:
        sid = stack.pop()
        s = id_to_step.get(sid)
        if not s:
            continue
        for dep in (s.get("depends_on") or []):
            if dep not in needed:
                needed.add(dep)
                stack.append(dep)

    # Build adjacency / indegree
    adj = {sid: [] for sid in needed}
    indeg = {sid: 0 for sid in needed}
    for sid in needed:
        s = id_to_step[sid]
        for dep in (s.get("depends_on") or []):
            if dep in needed:
                adj[dep].append(sid)
                indeg[sid] += 1

    file_map = dict(st.session_state.get("uploaded_file_paths", {}) or {})

    def write_paths_for(sid):
        s = id_to_step[sid]
        stype = (s.get("type") or "").lower().strip()
        cfg = s.get("config", {}) or {}
        labels = [x for x in _determine_modified_file_labels(stype, cfg) if x]
        return {str(file_map.get(lab)) for lab in labels if lab and file_map.get(lab)}

    node_paths = {sid: write_paths_for(sid) for sid in needed}

    progress = st.progress(0)
    status = st.empty()
    done = 0
    range_total = len(needed)
    results = []

    ready = {sid for sid, d in indeg.items() if d == 0}
    running = {}
    locked_paths = set()
    completed = set()

    max_workers = min(8, max(2, (os.cpu_count() or 4)))
    executor = cf.ThreadPoolExecutor(max_workers=max_workers)

    def run_step_worker(sid):
        step = id_to_step[sid]
        return execute_step_with_file_map(step, file_map)

    wb_cache.activate()
    try:
        while len(completed) < len(needed):
            # Run add_files / input_source in main thread (they mutate session_state)
            did_sync = True
            while did_sync:
                did_sync = False
                for sid in sorted(list(ready), key=lambda x: id_to_index.get(x, 10**9)):
                    stype = (id_to_step[sid].get("type") or "").lower().strip()
                    if stype in ("add_files", "input_source"):
                        ready.remove(sid)
                        ok, msg = execute_step(id_to_step[sid])
                        idx = id_to_index.get(sid, 0)
                        results.append({"step_num": idx + 1, "name": id_to_step[sid].get("name", ""), "type": stype, "success": ok, "message": msg})
                        if ok:
                            st.success(f"Step {idx + 1}: {msg}")
                            file_map.update(st.session_state.get("uploaded_file_paths", {}))
                            for nid in list(node_paths.keys()):
                                if nid not in completed:
                                    node_paths[nid] = write_paths_for(nid)
                        else:
                            st.error(f"Step {idx + 1}: {msg}")
                            return results
                        completed.add(sid)
                        done += 1
                        progress.progress(done / range_total)
                        for nxt in adj.get(sid, []):
                            indeg[nxt] -= 1
                            if indeg[nxt] == 0:
                                ready.add(nxt)
                        did_sync = True
                        break

            # Dispatch ready nodes
            dispatchable = []
            for sid in sorted(list(ready), key=lambda x: id_to_index.get(x, 10**9)):
                paths = node_paths.get(sid) or set()
                if paths and (paths & locked_paths):
                    continue
                dispatchable.append(sid)

            for sid in dispatchable:
                ready.remove(sid)
                paths = node_paths.get(sid) or set()
                locked_paths |= paths

                step = id_to_step[sid]
                swapped_paths, original_mappings = prepare_non_destructive_export(id_to_index.get(sid, 0) + 1, step)
                if swapped_paths:
                    file_map.update(st.session_state.get("uploaded_file_paths", {}))

                fut = executor.submit(run_step_worker, sid)
                running[fut] = (sid, paths, swapped_paths, original_mappings)

            if not running:
                if ready:
                    st.error("Dependency cycle detected.")
                else:
                    st.error("DAG could not make progress.")
                return results

            done_futs, _ = cf.wait(running.keys(), return_when=cf.FIRST_COMPLETED)
            for fut in done_futs:
                sid, paths, swapped_paths, original_mappings = running.pop(fut)
                locked_paths -= paths
                idx = id_to_index.get(sid, None)
                step = id_to_step[sid]
                try:
                    ok, msg = fut.result()
                except Exception as e:
                    ok, msg = False, str(e)

                if swapped_paths:
                    finalize_non_destructive_export(idx + 1 if isinstance(idx, int) else 0, swapped_paths, original_mappings)
                    file_map.update(st.session_state.get("uploaded_file_paths", {}))

                step_num = (idx + 1) if isinstance(idx, int) else "?"
                results.append({"step_num": step_num, "name": step.get("name", ""), "type": step.get("type", ""), "success": ok, "message": msg})

                if ok:
                    export_note = " (exported)" if swapped_paths else ""
                    st.success(f"Step {step_num}: {msg}{export_note}")
                else:
                    translated = translate_error(msg, step.get("type", ""))
                    st.error(f"Step {step_num} - {translated['title']}")
                    for other in list(running.keys()):
                        other.cancel()
                    return results

                completed.add(sid)
                done += 1
                progress.progress(done / range_total)

                for nxt in adj.get(sid, []):
                    indeg[nxt] -= 1
                    if indeg[nxt] == 0:
                        ready.add(nxt)

        status.text("Complete!")
        return results
    finally:
        wb_cache.deactivate()
        executor.shutdown(wait=False, cancel_futures=True)
