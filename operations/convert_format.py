"""
Convert Format operation.

Converts values in a selected column to a specific data type and applies the
matching Excel number format:

  • Number      — strip commas/currency symbols, parse to float; apply #,##0.00
  • Integer     — parse to whole number; apply 0
  • Text        — stringify; apply @ (text format)
  • Date        — parse flexibly (DD-MM-YYYY, YYYY-MM-DD, …); apply chosen date format
  • Percentage  — strip %, divide by 100 if needed, store as decimal; apply 0.00%

Scope options
  • single         — one named sheet
  • all_sheets     — every sheet in the workbook
  • selected_sheets — include / exclude filters

Column reference accepts: column letter (A, B, AA…), 1-based index (1, 2…),
or header name found in header_row.

Cells that cannot be converted are left unchanged (counted in skipped totals).
"""

from __future__ import annotations

import datetime
import os
import re
from typing import Optional

from core import wb_cache

# ---------------------------------------------------------------------------
# Format type registry
# ---------------------------------------------------------------------------

CONVERT_FORMAT_TYPES: dict[str, str] = {
    "number":     "Number  (decimal / float)",
    "integer":    "Integer  (whole number)",
    "text":       "Text  (string)",
    "date":       "Date",
    "percentage": "Percentage  (stores as 0.xx decimal)",
}
CONVERT_FORMAT_TYPE_KEYS:   list[str] = list(CONVERT_FORMAT_TYPES.keys())
CONVERT_FORMAT_TYPE_LABELS: list[str] = list(CONVERT_FORMAT_TYPES.values())

# Common date formats shown in the UI
DATE_FORMAT_OPTIONS: list[str] = [
    "DD-MM-YYYY",
    "DD/MM/YYYY",
    "YYYY-MM-DD",
    "MM/DD/YYYY",
    "MM-DD-YYYY",
    "DD MMM YYYY",
    "YYYY/MM/DD",
]

# Default Excel number-format codes
_DEFAULT_NUMBER_FORMAT: dict[str, str] = {
    "number":     "#,##0.00",
    "integer":    "0",
    "text":       "@",
    "date":       "DD-MM-YYYY",
    "percentage": "0.00%",
}

# Characters stripped before numeric parsing
_STRIP_RE = re.compile(r"[₹$£€,\s]")


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _resolve_col_idx(ws, col_ref: str, header_row: int) -> Optional[int]:
    """
    Resolve a column reference to a 1-based column index.
    Accepts: column letter (A, B, AA), integer string (1, 2),
             or header name found in header_row.
    """
    col_ref = str(col_ref).strip()
    if not col_ref:
        return None
    try:
        from openpyxl.utils import column_index_from_string
        return column_index_from_string(col_ref)
    except Exception:
        pass
    try:
        idx = int(col_ref)
        if idx >= 1:
            return idx
    except ValueError:
        pass
    col_lower = col_ref.lower()
    for cell in ws[header_row]:
        if cell.value is not None and str(cell.value).strip().lower() == col_lower:
            return cell.column
    return None


def _resolve_target_sheets(
    wb,
    scope: str,
    sheet_name: str,
    include_sheets: list,
    exclude_sheets: list,
) -> list[str]:
    """Return the list of sheet names to process according to scope."""
    all_sheets = wb.sheetnames
    if scope == "single":
        return [sheet_name] if sheet_name in all_sheets else []

    inc_lower = {s.strip().lower() for s in (include_sheets or []) if s.strip()}
    exc_lower = {s.strip().lower() for s in (exclude_sheets or []) if s.strip()}

    result = []
    for sn in all_sheets:
        sn_lower = sn.strip().lower()
        if exc_lower and sn_lower in exc_lower:
            continue
        if inc_lower and sn_lower not in inc_lower:
            continue
        result.append(sn)
    return result


def _clean_numeric(raw: str) -> str:
    """Strip common non-numeric decoration (commas, currency symbols, spaces)."""
    return _STRIP_RE.sub("", raw).strip()


def _convert_cell(cell, format_type: str, date_format: str, custom_number_format: str):
    """
    Attempt to convert cell.value to format_type.

    Returns (new_value, excel_number_format) or (None, None) if conversion fails.
    new_value = None means leave the cell unchanged.
    """
    raw = cell.value
    if raw is None:
        return None, None

    # ── Text ─────────────────────────────────────────────────────────────────
    if format_type == "text":
        return str(raw), custom_number_format or "@"

    raw_str = str(raw).strip()

    # ── Number ───────────────────────────────────────────────────────────────
    if format_type == "number":
        # Already numeric?
        if isinstance(raw, (int, float)) and not isinstance(raw, bool):
            val = float(raw)
        else:
            try:
                val = float(_clean_numeric(raw_str.rstrip("%")))
            except (ValueError, TypeError):
                return None, None
        return val, custom_number_format or "#,##0.00"

    # ── Integer ──────────────────────────────────────────────────────────────
    if format_type == "integer":
        if isinstance(raw, (int, float)) and not isinstance(raw, bool):
            val = int(raw)
        else:
            try:
                val = int(float(_clean_numeric(raw_str.rstrip("%"))))
            except (ValueError, TypeError):
                return None, None
        return val, custom_number_format or "0"

    # ── Percentage ───────────────────────────────────────────────────────────
    if format_type == "percentage":
        if isinstance(raw, (int, float)) and not isinstance(raw, bool):
            val = float(raw)
            if val > 1:
                val = val / 100.0
        else:
            try:
                cleaned = _clean_numeric(raw_str)
                val = float(cleaned)
                if raw_str.endswith("%") or val > 1:
                    val = val / 100.0
            except (ValueError, TypeError):
                return None, None
        return val, custom_number_format or "0.00%"

    # ── Date ─────────────────────────────────────────────────────────────────
    if format_type == "date":
        if isinstance(raw, (datetime.date, datetime.datetime)):
            dt = raw
        else:
            try:
                from dateutil import parser as _dp
                dt = _dp.parse(raw_str, dayfirst=True)
            except Exception:
                return None, None
        excel_fmt = custom_number_format or date_format or "DD-MM-YYYY"
        return dt, excel_fmt

    return None, None


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

def run_convert_format_step(
    file_path: str,
    sheet_name: str = "",
    scope: str = "single",               # "single" | "all_sheets" | "selected_sheets"
    include_sheets: list = None,
    exclude_sheets: list = None,
    column: str = "",                    # header name, letter, or 1-based index
    start_row: int = 2,                  # first data row (skip header)
    format_type: str = "number",
    date_format: str = "DD-MM-YYYY",     # Excel number-format string for dates
    custom_number_format: str = "",      # optional override for any type
    header_row: int = 1,
) -> tuple[bool, str]:
    """
    Convert values in *column* to *format_type* across one or more sheets.

    Parameters
    ----------
    file_path          : path to the workbook to modify
    sheet_name         : sheet name (used when scope == "single")
    scope              : "single" | "all_sheets" | "selected_sheets"
    include_sheets     : only process these sheet names (selected_sheets mode)
    exclude_sheets     : skip these sheet names
    column             : column to convert — header name, letter (A/B…), or 1-based int
    start_row          : first row to convert (1-based); default 2 skips the header
    format_type        : "number" | "integer" | "text" | "date" | "percentage"
    date_format        : Excel date format code (e.g. "DD-MM-YYYY") — only for date type
    custom_number_format : override Excel number-format code for any type
    header_row         : 1-based row that contains headers (for column name lookup)
    """

    # ── Validation ────────────────────────────────────────────────────────────
    if not file_path or not os.path.exists(file_path):
        return False, f"File not found: {file_path}"
    if not column:
        return False, "Column reference is required."
    if format_type not in CONVERT_FORMAT_TYPE_KEYS:
        return False, f"Unknown format type: '{format_type}'. Valid: {CONVERT_FORMAT_TYPE_KEYS}"

    # ── Load workbook ─────────────────────────────────────────────────────────
    try:
        wb = wb_cache.load(file_path)
    except Exception as exc:
        return False, f"Could not open file: {exc}"

    target_sheets = _resolve_target_sheets(
        wb, scope, sheet_name, include_sheets, exclude_sheets
    )
    if not target_sheets:
        return False, "No sheets matched the scope configuration."

    total_converted  = 0
    total_skipped    = 0
    col_not_found    = []
    sheets_touched   = 0

    for sn in target_sheets:
        ws = wb[sn]
        col_idx = _resolve_col_idx(ws, column, header_row)
        if col_idx is None:
            col_not_found.append(sn)
            continue

        converted_here = 0
        for row_idx in range(start_row, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value is None:
                continue
            new_val, num_fmt = _convert_cell(
                cell, format_type, date_format, custom_number_format
            )
            if new_val is None and num_fmt is None:
                total_skipped += 1
                continue
            cell.value = new_val
            if num_fmt:
                cell.number_format = num_fmt
            converted_here += 1

        if converted_here:
            sheets_touched += 1
        total_converted += converted_here

    # ── Save ──────────────────────────────────────────────────────────────────
    wb_cache.save(wb, file_path)

    # ── Summary ───────────────────────────────────────────────────────────────
    summary = (
        f"Converted {total_converted} cell(s) to {CONVERT_FORMAT_TYPES[format_type]} "
        f"across {sheets_touched} sheet(s)"
    )
    if total_skipped:
        summary += f". {total_skipped} cell(s) could not be converted and were left unchanged"
    if col_not_found:
        preview = col_not_found[:5]
        if len(col_not_found) > 5:
            preview.append(f"…+{len(col_not_found) - 5} more")
        summary += f". Column '{column}' not found in: {', '.join(preview)}"

    return True, summary
