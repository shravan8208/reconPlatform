"""
VLOOKUP operations using pandas merge.
Pure Python/pandas implementation - NO win32com!
"""

import pandas as pd
import openpyxl
from core import wb_cache

from typing import Optional, Tuple, List, Dict, Any

# Streamlit is optional - only used for warnings in UI mode
try:
    import streamlit as st
except ImportError:
    st = None


def vlookup_single(df_main, df_lookup, main_col, key_col, return_col, 
                   return_col_name=None, fallback_value=None, use_main_fallback=False,
                   fill_only_missing=False, fallback_column=None):
    """
    Perform a single VLOOKUP operation using pandas merge.
    
    Args:
        df_main: Main DataFrame
        df_lookup: Lookup DataFrame
        main_col: Column in main DataFrame to match on
        key_col: Column in lookup DataFrame to match on
        return_col: Column from lookup DataFrame to return
        return_col_name: Name for the returned column (default: same as return_col)
        fallback_value: Value to use if no match found
        use_main_fallback: Use main column value as fallback
        fill_only_missing: Only fill missing values
        fallback_column: Column from main DataFrame to use as fallback
    
    Returns:
        Modified DataFrame with lookup results
    """
    if df_main is None or df_lookup is None:
        return df_main
    
    if main_col not in df_main.columns:
        if st is not None:
            st.warning(f"Main column '{main_col}' not found in main data. Skipping.")
        return df_main
    
    if key_col not in df_lookup.columns or return_col not in df_lookup.columns:
        if st is not None:
            st.warning(f"Key or Return column not found in lookup data. Skipping.")
        return df_main
    
    # Get unique lookup values
    lookup_small = df_lookup[[key_col, return_col]].drop_duplicates(subset=[key_col])
    
    out_col = return_col_name if return_col_name else return_col
    
    # If fill_only_missing, preserve existing values
    if fill_only_missing and out_col in df_main.columns:
        missing_mask = df_main[out_col].isna() | (df_main[out_col].astype(str).str.strip() == '')
        if not missing_mask.any():
            return df_main
        
        df_to_lookup = df_main[missing_mask].copy()
        df_already_filled = df_main[~missing_mask].copy()
        
        merged = df_to_lookup.merge(
            lookup_small,
            how='left',
            left_on=main_col,
            right_on=key_col,
            suffixes=('', '_lookup')
        )
        
        # Combine back
        df_main = pd.concat([df_already_filled, merged], ignore_index=True)
    else:
        # Regular merge
        merged = df_main.merge(
            lookup_small,
            how='left',
            left_on=main_col,
            right_on=key_col,
            suffixes=('', '_lookup')
        )
        
        if return_col in merged.columns:
            merged[out_col] = merged[return_col]
        
        df_main = merged
    
    # Apply fallback logic
    if out_col in df_main.columns:
        missing = df_main[out_col].isna()
        
        if missing.any():
            if fallback_column and fallback_column in df_main.columns:
                df_main.loc[missing, out_col] = df_main.loc[missing, fallback_column]
            elif use_main_fallback and main_col in df_main.columns:
                df_main.loc[missing, out_col] = df_main.loc[missing, main_col]
            elif fallback_value is not None:
                df_main.loc[missing, out_col] = fallback_value
    
    return df_main


def run_vlookup_step(file_path, sheet_name, main_col, sources, fallback_value=None,
                     use_main_fallback=False, fill_only_missing=False, 
                     fallback_column=None, header_row=1):
    """
    Run VLOOKUP with multiple sources.
    
    Args:
        file_path: Path to main Excel file
        sheet_name: Sheet name
        main_col: Main column to match on
        sources: List of dicts with 'file_path', 'sheet_name', 'key_col', 'return_col'
        fallback_value: Default fallback value
        use_main_fallback: Use main column as fallback
        fill_only_missing: Only fill missing values
        fallback_column: Column to use as fallback
        header_row: Header row number (1-indexed)
    
    Returns:
        Tuple of (success: bool, message: str)
    """
    try:
        # Read main file
        df_main = wb_cache.read_excel(file_path, sheet_name=sheet_name, header=header_row-1)
        
        lookups_performed = 0
        
        # Perform each lookup
        for source in sources:
            lookup_file = source.get('file_path') or source.get('label')
            lookup_sheet = source.get('sheet_name') or source.get('sheet')
            key_col = source.get('key_col')
            return_col = source.get('return_col')
            return_col_name = source.get('return_col_name', return_col)
            
            if not lookup_file or not key_col or not return_col:
                continue
            
            # Read lookup file
            df_lookup = wb_cache.read_excel(lookup_file, sheet_name=lookup_sheet)
            
            # Perform lookup
            df_main = vlookup_single(
                df_main, df_lookup, main_col, key_col, return_col,
                return_col_name, fallback_value, use_main_fallback,
                fill_only_missing, fallback_column
            )
            
            lookups_performed += 1
        
        # Write back
        with wb_cache.excel_writer(file_path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
            df_main.to_excel(writer, sheet_name=sheet_name, index=False, startrow=header_row-1)
        
        return True, f"VLOOKUP completed, {lookups_performed} sources processed"
        
    except Exception as e:
        return False, f"Error in VLOOKUP: {str(e)}"


def run_vlookup_explicit_step(
    lookup_value_file_path: str,
    lookup_value_sheet_name: str,
    lookup_value_header_row: int,
    lookup_value_column,  # str or list of str (for composite key)
    search_file_path: str,
    search_sheet_name: str,
    search_header_row: int,
    search_column,  # str or list of str (multiple = search in ANY, OR logic)
    return_file_path: str,
    return_sheet_name: str,
    return_header_row: int,
    return_column: str,
    return_to_file_path: str,
    return_to_sheet_name: str,
    return_to_header_row: int,
    return_to_column: str,
    not_found_value: str = "",
    end_row: str = "last",
    fallback_lookups: Optional[List[Dict[str, Any]]] = None,
) -> Tuple[bool, str]:
    """
    VLOOKUP with explicit wiring of:
      1) Lookup Value (file/sheet/column(s))
      2) Search Value (file/sheet/column(s))
      3) Return Value (file/sheet/column) + not_found_value
      4) Return To (file/sheet/column) + return_to_header_row

    Matching modes:
    - 1 lookup col + N search cols: Search lookup value in ANY of the N columns (OR logic)
    - N lookup cols + N search cols: Composite key matching (AND logic)

    Notes:
    - Writes results into the *return_to* workbook in-place (preserves other rows/formatting).
    - Matches rows by relative data-row index.
    - 'end_row' can be 'last' or a string/int row number.
    """
    try:
        # Normalize columns to lists
        lookup_columns = [lookup_value_column] if isinstance(lookup_value_column, str) else list(lookup_value_column)
        search_columns = [search_column] if isinstance(search_column, str) else list(search_column)

        def _strip_cols(df: pd.DataFrame) -> pd.DataFrame:
            df = df.copy()
            df.columns = [str(c).strip() for c in df.columns]
            return df

        def _make_composite_key(df: pd.DataFrame, cols: List[str]) -> pd.Series:
            # Always strip each part to match how we build lookup_key from openpyxl
            parts = []
            for c in cols:
                c0 = str(c).strip()
                s = df[c0].astype(str).fillna("").map(lambda x: str(x).strip())
                parts.append(s)
            if len(parts) == 1:
                return parts[0]
            out = parts[0]
            for p in parts[1:]:
                out = out.str.cat(p, sep="|||")
            return out

        def _build_mapping(
            _search_file_path: str,
            _search_sheet_name: str,
            _search_header_row: int,
            _search_columns: List[str],
            _return_file_path: str,
            _return_sheet_name: str,
            _return_header_row: int,
            _return_column: str,
        ) -> tuple[Dict[str, Any], str]:
            """
            Build a mapping dict from (search key) -> return value.
            Matching mode depends on lookup_columns and the provided search columns.
            """
            df_s = wb_cache.read_excel(_search_file_path, sheet_name=_search_sheet_name, header=_search_header_row - 1)
            df_r = wb_cache.read_excel(_return_file_path, sheet_name=_return_sheet_name, header=_return_header_row - 1)
            df_s = _strip_cols(df_s)
            df_r = _strip_cols(df_r)

            for sc in _search_columns:
                sc0 = str(sc).strip()
                if sc0 not in df_s.columns:
                    return {}, f"❌ Search column '{sc}' not found. Available: {list(df_s.columns)[:10]}..."
            rc0 = str(_return_column).strip()
            if rc0 not in df_r.columns:
                return {}, f"❌ Return column '{_return_column}' not found"

            return_vals = df_r[rc0]

            # Determine matching mode
            if len(lookup_columns) == 1 and len(_search_columns) >= 1:
                # MODE: OR across N search columns
                mapping0: Dict[str, Any] = {}
                for sc in _search_columns:
                    sc0 = str(sc).strip()
                    for k, v in zip(df_s[sc0], return_vals):
                        if pd.isna(k) or str(k).strip() == "" or str(k) == "nan":
                            continue
                        key = str(k).strip()
                        if key not in mapping0:
                            mapping0[key] = v
                return mapping0, f"OR across {len(_search_columns)} col(s)"

            if len(lookup_columns) == len(_search_columns):
                # MODE: Composite key AND
                search_keys = _make_composite_key(df_s, _search_columns)
                mapping0 = {}
                for k, v in zip(search_keys, return_vals):
                    if pd.isna(k) or str(k).strip() == "" or str(k) == "nan":
                        continue
                    key = str(k).strip()
                    if key not in mapping0:
                        mapping0[key] = v
                return mapping0, f"composite key ({len(lookup_columns)} cols)"

            return {}, (
                f"❌ Invalid config: {len(lookup_columns)} lookup col(s) with {len(_search_columns)} search col(s). "
                f"Use 1 lookup + N search (OR), or N lookup + N search (AND)."
            )

        # Primary mapping
        mapping, mode_desc = _build_mapping(
            search_file_path,
            search_sheet_name,
            int(search_header_row),
            search_columns,
            return_file_path,
            return_sheet_name,
            int(return_header_row),
            return_column,
        )
        if not mapping and str(mode_desc).startswith("❌"):
            return False, str(mode_desc).replace("❌ ", "")

        # Fallback mappings
        fallback_mappings: List[tuple[Dict[str, Any], str]] = []
        for i, fb in enumerate(fallback_lookups or [], start=2):
            fb_search_columns = fb.get("search_columns")
            if fb_search_columns is None:
                fb_search_columns = fb.get("search_column")
            fb_search_cols_list = [fb_search_columns] if isinstance(fb_search_columns, str) else list(fb_search_columns or [])
            m, desc = _build_mapping(
                fb.get("search_file_path") or fb.get("search_file") or "",
                fb.get("search_sheet_name") or fb.get("search_sheet") or "",
                int(fb.get("search_header_row") or 1),
                fb_search_cols_list,
                fb.get("return_file_path") or fb.get("return_file") or "",
                fb.get("return_sheet_name") or fb.get("return_sheet") or "",
                int(fb.get("return_header_row") or 1),
                fb.get("return_column") or "",
            )
            if not m and str(desc).startswith("❌"):
                return False, f"Fallback VLOOKUP #{i}: {str(desc).replace('❌ ', '')}"
            fallback_mappings.append((m, desc))

        # Open destination workbook for writing
        wb_dest = wb_cache.load(return_to_file_path)
        ws_dest = wb_dest[return_to_sheet_name] if return_to_sheet_name in wb_dest.sheetnames else wb_dest.active

        # Open lookup-value workbook for reading
        same_lookup = (lookup_value_file_path == return_to_file_path)
        wb_lookup = wb_dest if same_lookup else wb_cache.load(lookup_value_file_path, data_only=True)
        ws_lookup = wb_lookup[lookup_value_sheet_name] if lookup_value_sheet_name in wb_lookup.sheetnames else wb_lookup.active

        def _find_col_idx(ws, header_row, header_name) -> Optional[int]:
            header_name_str = str(header_name).strip()
            if not header_name_str:
                return None
            max_col = ws.max_column
            for c in range(1, max_col + 1):
                v = ws.cell(row=int(header_row), column=c).value
                if v is None:
                    continue
                if str(v).strip() == header_name_str:
                    return c
            return None

        # Find column indices for all lookup columns
        lookup_col_idxs = []
        for lc in lookup_columns:
            idx = _find_col_idx(ws_lookup, lookup_value_header_row, lc)
            if idx is None:
                return False, f"Lookup Value column '{lc}' not found in header row {lookup_value_header_row}"
            lookup_col_idxs.append(idx)
        
        dest_col_idx = _find_col_idx(ws_dest, return_to_header_row, return_to_column)
        if dest_col_idx is None:
            return False, f"Return To column '{return_to_column}' not found in header row {return_to_header_row}"

        dest_first_data_row = int(return_to_header_row) + 1
        lookup_first_data_row = int(lookup_value_header_row) + 1

        if str(end_row).lower() == "last":
            dest_end_row = ws_dest.max_row
        else:
            dest_end_row = int(end_row)

        if dest_end_row < dest_first_data_row:
            return False, "End row is before the first data row"

        written = 0
        matched_primary = 0
        matched_fallback = [0 for _ in fallback_mappings]
        for dest_row in range(dest_first_data_row, dest_end_row + 1):
            i = dest_row - dest_first_data_row
            lookup_row = lookup_first_data_row + i
            
            # Build lookup key (single or composite)
            lookup_parts = []
            for col_idx in lookup_col_idxs:
                v = ws_lookup.cell(row=lookup_row, column=col_idx).value
                lookup_parts.append(str(v).strip() if v is not None else '')
            
            if len(lookup_parts) == 1:
                lookup_key = lookup_parts[0]
            else:
                lookup_key = '|||'.join(lookup_parts)

            out = mapping.get(lookup_key, None)
            found = not (out is None or (isinstance(out, float) and pd.isna(out)))
            if found:
                matched_primary += 1
            else:
                # Try fallback lookups in order
                for idx_fb, (m, _desc) in enumerate(fallback_mappings):
                    out2 = m.get(lookup_key, None)
                    found2 = not (out2 is None or (isinstance(out2, float) and pd.isna(out2)))
                    if found2:
                        out = out2
                        matched_fallback[idx_fb] += 1
                        found = True
                        break

            if not found:
                out = not_found_value

            ws_dest.cell(row=dest_row, column=dest_col_idx).value = out
            written += 1

        wb_cache.save(wb_dest, return_to_file_path)
        wb_dest.close()
        if not same_lookup:
            wb_lookup.close()

        col_desc = '+'.join(lookup_columns) if len(lookup_columns) > 1 else lookup_columns[0]
        parts = [f"VLOOKUP ({col_desc}, {mode_desc}) wrote {written} row(s), {matched_primary} matched"]
        for j, (_m, desc) in enumerate(fallback_mappings, start=2):
            parts.append(f"fallback #{j}: {matched_fallback[j-2]} matched ({desc})")
        parts.append(f"into '{return_to_column}'")
        return True, ", ".join(parts)

    except Exception as e:
        return False, f"Error in VLOOKUP: {str(e)}"

