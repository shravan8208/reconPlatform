"""
Advanced VLOOKUP with comprehensive condition builder and full dropdown interface.
Pure Python/pandas implementation - NO win32com!

Supports condition operators:
- equals, not_equals
- greater_than, less_than, greater_equal, less_equal
- starts_with, not_starts_with, ends_with, not_ends_with
- contains, not_contains
- is_blank, is_not_blank
"""

import pandas as pd
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows
from core import wb_cache
from openpyxl.utils import column_index_from_string, get_column_letter


def evaluate_condition(value, operator, compare_value):
    """
    Evaluate a condition using the specified operator.
    
    Args:
        value: The cell value to check
        operator: One of the condition operators
        compare_value: The value to compare against
    
    Returns:
        bool: True if condition is met
    """
    # Handle None/NaN
    if pd.isna(value):
        value = ""
    
    # Convert to string for string operations
    str_value = str(value).strip()
    str_compare = str(compare_value).strip() if compare_value else ""
    
    if operator == "equals":
        return str_value.upper() == str_compare.upper()
    
    elif operator == "not_equals":
        return str_value.upper() != str_compare.upper()
    
    elif operator == "greater_than":
        try:
            return float(value) > float(compare_value)
        except (ValueError, TypeError):
            return str_value > str_compare
    
    elif operator == "less_than":
        try:
            return float(value) < float(compare_value)
        except (ValueError, TypeError):
            return str_value < str_compare
    
    elif operator == "greater_equal":
        try:
            return float(value) >= float(compare_value)
        except (ValueError, TypeError):
            return str_value >= str_compare
    
    elif operator == "less_equal":
        try:
            return float(value) <= float(compare_value)
        except (ValueError, TypeError):
            return str_value <= str_compare
    
    elif operator == "starts_with":
        return str_value.upper().startswith(str_compare.upper())
    
    elif operator == "not_starts_with":
        return not str_value.upper().startswith(str_compare.upper())
    
    elif operator == "ends_with":
        return str_value.upper().endswith(str_compare.upper())
    
    elif operator == "not_ends_with":
        return not str_value.upper().endswith(str_compare.upper())
    
    elif operator == "contains":
        return str_compare.upper() in str_value.upper()
    
    elif operator == "not_contains":
        return str_compare.upper() not in str_value.upper()
    
    elif operator == "is_blank":
        return str_value == "" or pd.isna(value)
    
    elif operator == "is_not_blank":
        return str_value != "" and not pd.isna(value)
    
    # Default to equals
    return str_value.upper() == str_compare.upper()


def run_advance_vlookup_v2(config):
    """
    Advanced VLOOKUP with comprehensive condition builder.
    
    This is the new version that supports:
    - Full condition operators (equals, starts_with, contains, etc.)
    - Full VLOOKUP-style dropdown configuration
    - ELSE actions (copy column, fixed value, leave unchanged)
    
    Args:
        config: Dictionary with all configuration from the UI
    
    Returns:
        Tuple of (success: bool, message: str)
    """
    try:
        # Extract configuration
        condition_file = config.get("condition_file")
        condition_sheet = config.get("condition_sheet")
        condition_header_row = config.get("condition_header_row", 1)
        condition_column = config.get("condition_column")
        condition_operator = config.get("condition_operator", "equals")
        condition_value = config.get("condition_value", "")
        
        lv_file = config.get("lookup_value_file")
        lv_sheet = config.get("lookup_value_sheet")
        lv_header_row = config.get("lookup_value_header_row", 1)
        lv_columns = config.get("lookup_value_columns", [])
        
        s_file = config.get("search_file")
        s_sheet = config.get("search_sheet")
        s_header_row = config.get("search_header_row", 1)
        s_columns = config.get("search_columns", [])
        
        r_file = config.get("return_file")
        r_sheet = config.get("return_sheet")
        r_header_row = config.get("return_header_row", 1)
        r_column = config.get("return_column")
        not_found_value = config.get("not_found_value", "Not found")
        
        d_file = config.get("destination_file")
        d_sheet = config.get("destination_sheet")
        d_header_row = config.get("destination_header_row", 1)
        d_column = config.get("destination_column")
        end_row = config.get("end_row", "last")
        
        else_action = config.get("else_action", "Leave unchanged")
        else_source_file = config.get("else_source_file")
        else_source_sheet = config.get("else_source_sheet")
        else_source_column = config.get("else_source_column")
        else_fixed_value = config.get("else_fixed_value", "")
        
        # Read condition file (usually same as destination)
        df_cond = wb_cache.read_excel(condition_file, sheet_name=condition_sheet, header=condition_header_row-1)
        
        # Read lookup value file
        df_lookup_val = wb_cache.read_excel(lv_file, sheet_name=lv_sheet, header=lv_header_row-1)
        
        # Read search/return file (lookup table)
        df_search = wb_cache.read_excel(s_file, sheet_name=s_sheet, header=s_header_row-1)
        
        # Read return file (may be same as search)
        if r_file == s_file and r_sheet == s_sheet:
            df_return = df_search
        else:
            df_return = wb_cache.read_excel(r_file, sheet_name=r_sheet, header=r_header_row-1)
        
        # Read destination file
        df_dest = wb_cache.read_excel(d_file, sheet_name=d_sheet, header=d_header_row-1)
        
        # Read else source file if needed
        df_else = None
        if else_action == "Copy from another column" and else_source_file:
            if else_source_file == d_file and else_source_sheet == d_sheet:
                df_else = df_dest
            else:
                df_else = wb_cache.read_excel(else_source_file, sheet_name=else_source_sheet, header=d_header_row-1)
        
        # Build lookup key from lookup value columns
        if len(lv_columns) == 1:
            # Single column lookup
            if lv_columns[0] not in df_lookup_val.columns:
                return False, f"Lookup column '{lv_columns[0]}' not found"
            df_lookup_val['_lookup_key'] = df_lookup_val[lv_columns[0]].astype(str).str.strip()
        else:
            # Composite key (concatenate multiple columns)
            for col in lv_columns:
                if col not in df_lookup_val.columns:
                    return False, f"Lookup column '{col}' not found"
            df_lookup_val['_lookup_key'] = df_lookup_val[lv_columns].astype(str).apply(
                lambda x: '|||'.join(x.str.strip()), axis=1
            )
        
        # Build search key from search columns
        if len(s_columns) == 1 and len(lv_columns) == 1:
            # Single column search - OR mode if needed
            if s_columns[0] not in df_search.columns:
                return False, f"Search column '{s_columns[0]}' not found"
            df_search['_search_key'] = df_search[s_columns[0]].astype(str).str.strip()
        elif len(lv_columns) == 1 and len(s_columns) > 1:
            # OR mode: single lookup value, multiple search columns
            # Build dict for each search column
            pass  # Will handle in lookup phase
        else:
            # AND mode: composite key
            for col in s_columns:
                if col not in df_search.columns:
                    return False, f"Search column '{col}' not found"
            df_search['_search_key'] = df_search[s_columns].astype(str).apply(
                lambda x: '|||'.join(x.str.strip()), axis=1
            )
        
        # Return column validation
        if r_column not in df_return.columns:
            return False, f"Return column '{r_column}' not found"
        
        # Build lookup dictionary
        if len(lv_columns) == 1 and len(s_columns) > 1:
            # OR mode: check against multiple search columns
            lookup_dicts = []
            for s_col in s_columns:
                if s_col in df_search.columns:
                    lookup_dicts.append(dict(zip(
                        df_search[s_col].astype(str).str.strip(),
                        df_return[r_column]
                    )))
        else:
            # Standard or AND mode
            lookup_dict = dict(zip(df_search['_search_key'], df_return[r_column]))
        
        # Initialize output column
        if d_column not in df_dest.columns:
            df_dest[d_column] = None
        
        # Find condition column
        if condition_column in df_cond.columns:
            cond_col_name = condition_column
        else:
            # Try as column letter
            try:
                col_idx = column_index_from_string(condition_column) - 1
                cond_col_name = df_cond.columns[col_idx]
            except:
                return False, f"Condition column '{condition_column}' not found"
        
        # Evaluate condition for each row and perform VLOOKUP or ELSE
        vlookup_count = 0
        else_count = 0
        not_found_count = 0
        
        for idx in range(len(df_dest)):
            # Get condition value
            cond_val = df_cond.iloc[idx][cond_col_name] if idx < len(df_cond) else None
            
            # Evaluate condition
            condition_met = evaluate_condition(cond_val, condition_operator, condition_value)
            
            if condition_met:
                # Do VLOOKUP
                lookup_key = df_lookup_val.iloc[idx]['_lookup_key'] if idx < len(df_lookup_val) else None
                
                if lookup_key:
                    result = None
                    
                    if len(lv_columns) == 1 and len(s_columns) > 1:
                        # OR mode: check multiple dicts
                        for ld in lookup_dicts:
                            if lookup_key in ld:
                                result = ld[lookup_key]
                                break
                    else:
                        # Standard lookup
                        result = lookup_dict.get(lookup_key)
                    
                    if result is not None:
                        df_dest.at[idx, d_column] = result
                        vlookup_count += 1
                    else:
                        df_dest.at[idx, d_column] = not_found_value
                        not_found_count += 1
                else:
                    df_dest.at[idx, d_column] = not_found_value
                    not_found_count += 1
            else:
                # ELSE action
                if else_action == "Copy from another column" and df_else is not None and else_source_column:
                    if else_source_column in df_else.columns and idx < len(df_else):
                        df_dest.at[idx, d_column] = df_else.iloc[idx][else_source_column]
                elif else_action == "Write a fixed value":
                    df_dest.at[idx, d_column] = else_fixed_value
                # "Leave unchanged" - do nothing
                
                else_count += 1
        
        # Write back to destination file
        wb = wb_cache.load(d_file)
        if d_sheet in wb.sheetnames:
            del wb[d_sheet]
        ws = wb.create_sheet(d_sheet)
        
        for r_idx, row in enumerate(dataframe_to_rows(df_dest, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                ws.cell(row=r_idx + d_header_row - 1, column=c_idx, value=value)
        
        wb_cache.save(wb, d_file)
        wb.close()
        
        # Build result message
        op_names = {
            "equals": "=", "not_equals": "<>", "greater_than": ">", "less_than": "<",
            "greater_equal": ">=", "less_equal": "<=", "starts_with": "starts with",
            "not_starts_with": "doesn't start with", "ends_with": "ends with",
            "not_ends_with": "doesn't end with", "contains": "contains",
            "not_contains": "doesn't contain", "is_blank": "is blank", "is_not_blank": "is not blank"
        }
        op_display = op_names.get(condition_operator, condition_operator)
        
        msg = f"Adv VLOOKUP: {vlookup_count} rows matched (IF {condition_column} {op_display}"
        if condition_operator not in ["is_blank", "is_not_blank"]:
            msg += f" '{condition_value}'"
        msg += f"), {else_count} rows ELSE"
        if not_found_count > 0:
            msg += f", {not_found_count} not found"
        
        return True, msg
        
    except Exception as e:
        return False, f"Error in advanced VLOOKUP: {str(e)}"


def run_advance_vlookup_step(file_path, sheet_name, name_prefix, source_column,
                             lookup_file_path, lookup_sheet_name, key_col, return_col,
                             output_col, fallback_value=None, fallback_col=None,
                             target_header_row=1, lookup_header_row=1,
                             # Conditional parameters (legacy)
                             condition_enabled=False,
                             condition_column=None,
                             condition_value=None,
                             else_source_column=None,
                             ):
    """
    Legacy Advanced VLOOKUP that concatenates a prefix with source column value as the key.
    Kept for backward compatibility with existing workflows.
    """
    try:
        # Read target file
        df_target = wb_cache.read_excel(file_path, sheet_name=sheet_name, header=target_header_row-1)
        
        # Read lookup file
        df_lookup = wb_cache.read_excel(lookup_file_path, sheet_name=lookup_sheet_name, header=lookup_header_row-1)
        
        # Create concatenated key in target DataFrame
        if source_column not in df_target.columns:
            return False, f"Source column '{source_column}' not found in target file"
        
        df_target['_lookup_key'] = name_prefix + df_target[source_column].astype(str)
        
        # Perform lookup
        if key_col not in df_lookup.columns or return_col not in df_lookup.columns:
            return False, f"Key column '{key_col}' or return column '{return_col}' not found in lookup file"
        
        lookup_dict = dict(zip(df_lookup[key_col], df_lookup[return_col]))
        
        # Initialize output column
        if output_col not in df_target.columns:
            df_target[output_col] = None
        
        # Apply conditional logic if enabled
        if condition_enabled and condition_column:
            # Find condition column
            cond_col_name = None
            if condition_column in df_target.columns:
                cond_col_name = condition_column
            else:
                try:
                    col_idx = column_index_from_string(condition_column) - 1
                    if col_idx < len(df_target.columns):
                        cond_col_name = df_target.columns[col_idx]
                except:
                    pass
            
            if not cond_col_name:
                return False, f"Condition column '{condition_column}' not found"
            
            # Find else source column
            else_col_name = None
            if else_source_column:
                if else_source_column in df_target.columns:
                    else_col_name = else_source_column
                else:
                    try:
                        col_idx = column_index_from_string(else_source_column) - 1
                        if col_idx < len(df_target.columns):
                            else_col_name = df_target.columns[col_idx]
                    except:
                        pass
            
            # Create condition mask
            condition_mask = df_target[cond_col_name].astype(str).str.strip().str.upper() == str(condition_value).strip().upper()
            
            # For rows where condition IS met: do VLOOKUP
            df_target.loc[condition_mask, output_col] = df_target.loc[condition_mask, '_lookup_key'].map(lookup_dict)
            
            # Apply fallback for VLOOKUP misses
            vlookup_missing = condition_mask & df_target[output_col].isna()
            if vlookup_missing.any():
                if fallback_col and fallback_col in df_target.columns:
                    df_target.loc[vlookup_missing, output_col] = df_target.loc[vlookup_missing, fallback_col]
                elif fallback_value is not None:
                    df_target.loc[vlookup_missing, output_col] = fallback_value
            
            # For rows where condition is NOT met: copy from else_source_column
            if else_col_name:
                df_target.loc[~condition_mask, output_col] = df_target.loc[~condition_mask, else_col_name]
            
            vlookup_count = condition_mask.sum()
            else_count = (~condition_mask).sum()
            action_msg = f"Conditional VLOOKUP: {vlookup_count} rows matched '{condition_value}' (VLOOKUP), {else_count} rows copied from {else_source_column}"
        
        else:
            # Standard VLOOKUP (no condition)
            df_target[output_col] = df_target['_lookup_key'].map(lookup_dict)
            
            # Apply fallback
            missing_mask = df_target[output_col].isna()
            if missing_mask.any():
                if fallback_col and fallback_col in df_target.columns:
                    df_target.loc[missing_mask, output_col] = df_target.loc[missing_mask, fallback_col]
                elif fallback_value is not None:
                    df_target.loc[missing_mask, output_col] = fallback_value
            
            action_msg = f"Advanced VLOOKUP completed: {name_prefix} + {source_column} → {output_col}"
        
        # Remove helper column
        df_target = df_target.drop(columns=['_lookup_key'])
        
        # Write back
        wb = wb_cache.load(file_path)
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(sheet_name)
        
        for r_idx, row in enumerate(dataframe_to_rows(df_target, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                ws.cell(row=r_idx + target_header_row - 1, column=c_idx, value=value)
        
        wb_cache.save(wb, file_path)
        wb.close()
        
        return True, action_msg
        
    except Exception as e:
        return False, f"Error in advanced VLOOKUP: {str(e)}"


