"""
Generic, reusable operations (no win32com).

These are designed to avoid creating one-off functions for every workflow:

1) quota_update:
   - Build per-key quotas from a filtered "source" table
   - Scan a "target" table and apply a capped update per key

2) group_rollup:
   - Roll up values across row types within groups (keys)
   - Optionally delete duplicate rows after consolidation
"""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
import pandas as pd
from openpyxl.utils import column_index_from_string
from core import wb_cache


# -----------------------------------------------------------------------------
# Helpers
# -----------------------------------------------------------------------------

def _try_float(v: Any) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)) and not pd.isna(v):
        return float(v)
    s = str(v).strip()
    if not s:
        return None
    try:
        return float(s)
    except Exception:
        return None


def _norm_value(v: Any, mode: str) -> str:
    """
    Normalize a value to a comparable string representation.

    Modes:
    - "text": string strip
    - "upper_text": string strip + upper()
    - "int_string": numeric -> integer-like string when possible (100.0 -> "100")
    """
    m = (mode or "text").strip().lower()
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    s = str(v).strip()

    if m == "upper_text":
        return s.upper()

    if m == "int_string":
        f = _try_float(v)
        if f is None:
            return s
        if f.is_integer():
            return str(int(f))
        # keep as trimmed float string (avoid scientific where possible)
        return str(f)

    # default: text
    return s


def _apply_filter(series: pd.Series, operator: str, value: str) -> pd.Series:
    op = (operator or "equals").strip().lower()
    v = "" if value is None else str(value)

    if op == "equals":
        return series.astype(str).str.strip() == v.strip()
    if op == "not_equals":
        return series.astype(str).str.strip() != v.strip()
    if op == "contains":
        return series.astype(str).str.contains(v, case=False, na=False)
    if op == "starts_with":
        return series.astype(str).str.startswith(v, na=False)
    if op == "ends_with":
        return series.astype(str).str.endswith(v, na=False)

    # Fallback to equals
    return series.astype(str).str.strip() == v.strip()


def _find_col_idx_by_header(ws, header_row: int, header_name: str) -> Optional[int]:
    name = str(header_name or "").strip()
    if not name:
        return None

    # Prefer treating as column letter only when plausible
    if name.isalpha() and 1 <= len(name) <= 3:
        try:
            idx = column_index_from_string(name.upper())
            if idx <= ws.max_column:
                return idx
        except Exception:
            pass

    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=int(header_row), column=c).value
        if v is None:
            continue
        if str(v).strip() == name:
            return c
    return None


# -----------------------------------------------------------------------------
# 1) quota_update
# -----------------------------------------------------------------------------

def run_quota_update_step(
    source_path: str,
    source_sheet: str,
    source_header_row: int,
    source_filter_column: str,
    source_filter_operator: str,
    source_filter_value: str,
    source_key_columns: List[str],
    source_key_norms: Optional[List[str]],
    target_path: str,
    target_sheet: str,
    target_header_row: int,
    target_key_columns: List[str],
    target_key_norms: Optional[List[str]],
    write_column: str,
    write_value: str,
    end_row: str = "last",
) -> Tuple[bool, str]:
    """
    Generic quota-based update:
    - Filter SOURCE rows and count frequency per source_key
    - Scan TARGET rows; for each row, if its target_key exists in allowed counts and
      hasn't reached its quota, write write_value into write_column.
    """
    try:
        if not source_path or not target_path:
            return False, "Missing source/target file path"
        if not source_sheet or not target_sheet:
            return False, "Missing source/target sheet"
        if not source_filter_column:
            return False, "Missing source filter column"
        if not source_key_columns or not target_key_columns:
            return False, "Missing key columns"
        if len(source_key_columns) != len(target_key_columns):
            return False, "Source key columns count must equal target key columns count"

        src_norms = list(source_key_norms or [])
        tgt_norms = list(target_key_norms or [])
        while len(src_norms) < len(source_key_columns):
            src_norms.append("text")
        while len(tgt_norms) < len(target_key_columns):
            tgt_norms.append("text")

        # Build allowed counts from SOURCE using pandas
        df = wb_cache.read_excel(source_path, sheet_name=source_sheet, header=int(source_header_row) - 1)
        df.columns = [str(c).strip() for c in df.columns]

        if source_filter_column not in df.columns:
            return False, f"Source filter column not found: {source_filter_column}"
        for c in source_key_columns:
            if str(c).strip() not in df.columns:
                return False, f"Source key column not found: {c}"

        mask = _apply_filter(df[source_filter_column], source_filter_operator, source_filter_value)
        df_f = df[mask].copy()
        if df_f.empty:
            return True, "Quota update: 0 source rows matched filter; nothing to apply"

        # Normalize source keys
        key_tuples = []
        for _, row in df_f.iterrows():
            parts = []
            for col, nm in zip(source_key_columns, src_norms):
                parts.append(_norm_value(row[str(col).strip()], nm))
            key_tuples.append(tuple(parts))

        allowed = Counter(key_tuples)

        # Apply to TARGET via openpyxl in-place
        wb = wb_cache.load(target_path)
        ws = wb[target_sheet] if target_sheet in wb.sheetnames else wb.active

        # Resolve header -> col idx
        key_idxs = []
        for c in target_key_columns:
            idx = _find_col_idx_by_header(ws, int(target_header_row), c)
            if idx is None:
                wb.close()
                return False, f"Target key column not found: {c} (header row {target_header_row})"
            key_idxs.append(idx)

        write_idx = _find_col_idx_by_header(ws, int(target_header_row), write_column)
        if write_idx is None:
            wb.close()
            return False, f"Target write column not found: {write_column} (header row {target_header_row})"

        first_row = int(target_header_row) + 1
        if str(end_row).lower() == "last" or end_row == "":
            last_row = ws.max_row
        else:
            last_row = int(end_row)
        if last_row < first_row:
            wb.close()
            return False, "End row is before first data row"

        marked = defaultdict(int)
        written = 0
        for r in range(first_row, last_row + 1):
            parts = []
            for idx, nm in zip(key_idxs, tgt_norms):
                v = ws.cell(row=r, column=idx).value
                parts.append(_norm_value(v, nm))
            k = tuple(parts)
            quota = allowed.get(k, 0)
            if quota <= 0:
                continue
            if marked[k] >= quota:
                continue
            ws.cell(row=r, column=write_idx).value = write_value
            marked[k] += 1
            written += 1

        wb_cache.save(wb, target_path)
        wb.close()

        return True, f"Quota update: wrote {written} row(s) into '{write_column}' (source matched {len(df_f)} rows; {len(allowed)} unique key(s))"
    except Exception as e:
        return False, f"Error in quota update: {str(e)}"


# -----------------------------------------------------------------------------
# 2) group_rollup
# -----------------------------------------------------------------------------

def run_group_rollup_step(
    file_path: str,
    sheet_name: str,
    header_row: int,
    key_columns: List[str],
    type_column: str,
    amount_column: str,
    from_type_value: str,
    into_type_value: str,
    include_from_in_sum: bool = True,
    delete_duplicate_into_rows: bool = True,
    delete_from_rows: bool = False,
    end_row: str = "last",
) -> Tuple[bool, str]:
    """
    Generic rollup + de-dup:
    For each key:
      - Identify rows where type == from_type_value and type == into_type_value
      - If key has at least one from-row and at least one into-row:
          merged = sum(into amounts) + (sum(from amounts) if include_from_in_sum)
          write merged into FIRST into-row amount cell
          delete remaining into-rows (optional)
          optionally delete from-rows (optional)
    """
    try:
        if not file_path or not sheet_name:
            return False, "Missing file or sheet"
        if not key_columns:
            return False, "Missing key columns"

        wb = wb_cache.load(file_path)
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active

        hdr = int(header_row)
        first_row = hdr + 1
        if str(end_row).lower() == "last" or end_row == "":
            last_row = ws.max_row
        else:
            last_row = int(end_row)
        if last_row < first_row:
            wb.close()
            return False, "End row is before first data row"

        key_idxs = []
        for c in key_columns:
            idx = _find_col_idx_by_header(ws, hdr, c)
            if idx is None:
                wb.close()
                return False, f"Key column not found: {c} (header row {header_row})"
            key_idxs.append(idx)

        type_idx = _find_col_idx_by_header(ws, hdr, type_column)
        if type_idx is None:
            wb.close()
            return False, f"Type column not found: {type_column} (header row {header_row})"

        amt_idx = _find_col_idx_by_header(ws, hdr, amount_column)
        if amt_idx is None:
            wb.close()
            return False, f"Amount column not found: {amount_column} (header row {header_row})"

        def make_key(r: int) -> Tuple[str, ...]:
            parts = []
            for idx in key_idxs:
                parts.append(_norm_value(ws.cell(row=r, column=idx).value, "text"))
            return tuple(parts)

        from_val = str(from_type_value or "").strip()
        into_val = str(into_type_value or "").strip()

        groups: Dict[Tuple[str, ...], Dict[str, List[int]]] = {}
        for r in range(first_row, last_row + 1):
            k = make_key(r)
            t = ws.cell(row=r, column=type_idx).value
            t_s = str(t).strip() if t is not None else ""
            if k not in groups:
                groups[k] = {"from": [], "into": []}
            if t_s == from_val:
                groups[k]["from"].append(r)
            elif t_s == into_val:
                groups[k]["into"].append(r)

        updated_keys = 0
        deleted_rows: List[int] = []

        for k, info in groups.items():
            from_rows = info["from"]
            into_rows = info["into"]
            if not from_rows or not into_rows:
                continue

            # Sum into amounts
            into_sum = 0.0
            for r in into_rows:
                v = ws.cell(row=r, column=amt_idx).value
                f = _try_float(v)
                into_sum += (f if f is not None else 0.0)

            from_sum = 0.0
            if include_from_in_sum:
                for r in from_rows:
                    v = ws.cell(row=r, column=amt_idx).value
                    f = _try_float(v)
                    from_sum += (f if f is not None else 0.0)

            merged = into_sum + from_sum

            # Write to first into row
            target_row = into_rows[0]
            # If it's integer-like, write as int to match Excel expectations
            if float(merged).is_integer():
                ws.cell(row=target_row, column=amt_idx).value = int(merged)
            else:
                ws.cell(row=target_row, column=amt_idx).value = float(merged)

            # Delete remaining into rows
            if delete_duplicate_into_rows and len(into_rows) > 1:
                deleted_rows.extend(into_rows[1:])

            # Optionally delete from rows
            if delete_from_rows:
                deleted_rows.extend(from_rows)

            updated_keys += 1

        # Delete rows bottom-up
        if deleted_rows:
            for r in sorted(set(deleted_rows), reverse=True):
                ws.delete_rows(r, 1)

        wb_cache.save(wb, file_path)
        wb.close()

        return True, f"Group rollup: updated {updated_keys} key(s); deleted {len(set(deleted_rows))} row(s)"
    except Exception as e:
        return False, f"Error in group rollup: {str(e)}"

