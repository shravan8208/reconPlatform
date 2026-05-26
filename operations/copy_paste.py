"""
Copy/Paste operations for Excel files.
Pure openpyxl implementation - NO win32com!
"""

import openpyxl
from openpyxl.utils import column_index_from_string
from typing import Iterable, List, Tuple, Union
from core import wb_cache


def _resolve_col_idx(ws, header_row: int, col_or_header: str) -> Union[int, None]:
    """
    Resolve a column identifier to a 1-based column index.

    Accepts:
    - Excel column letter(s): "A", "AB"
    - Header name (string) present in header_row
    """
    s = str(col_or_header or "").strip()
    if not s:
        return None

    # Prefer treating as column letter only when it plausibly fits sheet width.
    if s.isalpha() and 1 <= len(s) <= 3:
        try:
            idx = column_index_from_string(s.upper())
            if idx <= ws.max_column:
                return idx
        except Exception:
            pass

    # Otherwise, search headers for exact match (trimmed)
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=int(header_row), column=c).value
        if v is None:
            continue
        if str(v).strip() == s:
            return c
    return None


def run_copy_paste_step(
    file_path: str,
    sheet_name: str,
    src_col: str,
    tgt_col: Union[str, List[str]],
    paste_type: str,
    header_row: int = 1,
) -> Tuple[bool, str]:
    """
    Copy data from a source column to one OR multiple target columns.
    
    Args:
        file_path: Path to Excel file
        sheet_name: Sheet name
        src_col: Source column letter (e.g., 'A') OR header name (from header_row)
        tgt_col: Target column letter/header OR list of target columns/headers
        paste_type: 'All' (formulas + values) or 'Values Only'
        header_row: Row number where headers are (default: 1)
    
    Returns:
        Tuple of (success: bool, message: str)
    """
    if not src_col or not tgt_col:
        return False, "Missing Source or Target Column"
    
    try:
        # Determine if we need values only
        values_only = 'values' in paste_type.lower() if isinstance(paste_type, str) else False
        
        # Load workbook
        wb = wb_cache.load(file_path)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

        # Normalize targets to a list
        if isinstance(tgt_col, str):
            tgt_cols = [tgt_col]
        else:
            tgt_cols = list(tgt_col)
        tgt_cols = [str(c).strip() for c in tgt_cols if str(c).strip()]
        if not tgt_cols:
            return False, "Missing Target Column(s)"

        src_idx = _resolve_col_idx(ws, int(header_row), str(src_col))
        if src_idx is None:
            return False, f"Source column '{src_col}' not found (as letter or header in row {header_row})"

        tgt_idxs: List[int] = []
        for t in tgt_cols:
            idx = _resolve_col_idx(ws, int(header_row), t)
            if idx is None:
                return False, f"Target column '{t}' not found (as letter or header in row {header_row})"
            if idx not in tgt_idxs:
                tgt_idxs.append(idx)
        
        max_row = ws.max_row
        
        if values_only:
            # Load with data_only=True to get calculated values
            wb_data = wb_cache.load(file_path, data_only=True)
            ws_data = wb_data[sheet_name] if sheet_name and sheet_name in wb_data.sheetnames else wb_data.active
            
            # Copy values only (skip header row)
            for row in range(header_row + 1, max_row + 1):
                src_cell = ws_data.cell(row=row, column=src_idx)
                for tgt_idx in tgt_idxs:
                    ws.cell(row=row, column=tgt_idx).value = src_cell.value
            
            wb_data.close()
        else:
            # Copy everything including formulas (skip header row)
            for row in range(header_row + 1, max_row + 1):
                src_cell = ws.cell(row=row, column=src_idx)
                for tgt_idx in tgt_idxs:
                    tgt_cell = ws.cell(row=row, column=tgt_idx)

                    # Copy value (formula or data)
                    tgt_cell.value = src_cell.value

                    # Copy formatting
                    if src_cell.has_style:
                        tgt_cell.font = src_cell.font.copy()
                        tgt_cell.border = src_cell.border.copy()
                        tgt_cell.fill = src_cell.fill.copy()
                        tgt_cell.number_format = src_cell.number_format
                        tgt_cell.protection = src_cell.protection.copy()
                        tgt_cell.alignment = src_cell.alignment.copy()
        
        wb_cache.save(wb, file_path)
        wb.close()
        
        paste_mode = "values only" if values_only else "all (content + formulas)"
        tgt_desc = ", ".join(tgt_cols)
        return True, f"Copied {src_col} to {tgt_desc} ({paste_mode}), {max_row - header_row} rows"
        
    except Exception as e:
        return False, f"Error in copy/paste: {str(e)}"


