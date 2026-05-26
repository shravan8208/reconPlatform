"""
Merge cells operations.
Pure openpyxl implementation - NO win32com!
"""

import openpyxl
from core import wb_cache


def run_merge_cells_step(file_path, sheet_name, merge_range):
    """
    Merge a range of cells.
    
    Args:
        file_path: Path to Excel file
        sheet_name: Sheet name
        merge_range: Range to merge (e.g., "A1:C1")
    
    Returns:
        Tuple of (success: bool, message: str)
    """
    if not merge_range or ':' not in merge_range:
        return False, "Invalid merge range (must be like 'A1:C1')"
    
    try:
        wb = wb_cache.load(file_path)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        
        # Merge cells
        ws.merge_cells(merge_range)
        
        wb_cache.save(wb, file_path)
        wb.close()
        
        return True, f"Merged cells {merge_range}"
        
    except Exception as e:
        return False, f"Error merging cells: {str(e)}"


def run_unmerge_cells_step(file_path, sheet_name, merge_range):
    """
    Unmerge a range of cells.
    
    Args:
        file_path: Path to Excel file
        sheet_name: Sheet name
        merge_range: Range to unmerge (e.g., "A1:C1")
    
    Returns:
        Tuple of (success: bool, message: str)
    """
    if not merge_range or ':' not in merge_range:
        return False, "Invalid merge range (must be like 'A1:C1')"
    
    try:
        wb = wb_cache.load(file_path)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        
        # Unmerge cells
        ws.unmerge_cells(merge_range)
        
        wb_cache.save(wb, file_path)
        wb.close()
        
        return True, f"Unmerged cells {merge_range}"
        
    except Exception as e:
        return False, f"Error unmerging cells: {str(e)}"


