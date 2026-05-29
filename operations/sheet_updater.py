"""
Sheet Updater operation.

Reads a master Excel sheet with three key columns:
  • Category   — used to match a target sheet by name
  • Particulars — used to find the row inside that sheet (lookup)
  • Value       — the value to write into a specified column of that row

For every row in the master the engine:
  1. Finds the target sheet whose name matches the category value
     (configurable match mode: contains / starts_with / ends_with / exact)
  2. Searches the lookup column in that sheet for the particulars value
  3. Writes the value into the write column of the matching row

All reads/writes go through wb_cache so the workbook is only touched
once per workflow run.
"""

from __future__ import annotations

import os
from typing import Optional

import pandas as pd
import openpyxl

from core import wb_cache

# ---------------------------------------------------------------------------
# Sheet-match mode registry
# ---------------------------------------------------------------------------

SHEET_MATCH_MODES: dict[str, str] = {
    "cat_in_sheet": "Category value is contained in sheet name",
    "sheet_in_cat": "Sheet name is contained in category value",
    "starts_with":  "Sheet name starts with category value",
    "ends_with":    "Sheet name ends with category value",
    "exact":        "Exact match (case-insensitive)",
}
SHEET_MATCH_KEYS:   list[str] = list(SHEET_MATCH_MODES.keys())
SHEET_MATCH_LABELS: list[str] = list(SHEET_MATCH_MODES.values())


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _match_sheet(category: str, sheet_names: list[str], mode: str) -> Optional[str]:
    """Return the first sheet whose name satisfies the match condition."""
    cat = str(category).strip().lower()
    for sn in sheet_names:
        sn_lower = sn.strip().lower()
        if mode == "exact":
            if cat == sn_lower:
                return sn
        elif mode == "starts_with":
            if sn_lower.startswith(cat):
                return sn
        elif mode == "ends_with":
            if sn_lower.endswith(cat):
                return sn
        elif mode == "cat_in_sheet":
            if cat in sn_lower:
                return sn
        elif mode == "sheet_in_cat":
            if sn_lower in cat:
                return sn
    return None


def _resolve_col_idx(ws, col_ref: str) -> Optional[int]:
    """
    Resolve a column reference to a 1-based column index.
    Accepts: column letter ('A', 'B', 'AA'), integer string ('1', '3'),
             or header name found in row 1.
    """
    col_ref = str(col_ref).strip()

    # Column letter (A, B, AA …)
    try:
        from openpyxl.utils import column_index_from_string
        return column_index_from_string(col_ref)
    except Exception:
        pass

    # Integer index
    try:
        idx = int(col_ref)
        if idx >= 1:
            return idx
    except ValueError:
        pass

    # Header name — scan row 1
    col_lower = col_ref.lower()
    for cell in ws[1]:
        if cell.value is not None and str(cell.value).strip().lower() == col_lower:
            return cell.column

    return None


def _find_df_col(df: pd.DataFrame, ref: str) -> Optional[str]:
    """
    Resolve a column reference (header name, letter, or 1-based integer)
    to the actual DataFrame column name.
    """
    ref = str(ref).strip()
    ref_lower = ref.lower()

    # Exact header name match
    for c in df.columns:
        if str(c).strip().lower() == ref_lower:
            return c

    # Column letter → 0-based index
    try:
        from openpyxl.utils import column_index_from_string
        idx = column_index_from_string(ref) - 1
        if 0 <= idx < len(df.columns):
            return df.columns[idx]
    except Exception:
        pass

    # Integer → 0-based index
    try:
        idx = int(ref) - 1
        if 0 <= idx < len(df.columns):
            return df.columns[idx]
    except ValueError:
        pass

    return None


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

def run_sheet_updater_step(
    src_file_path: str,
    src_sheet: str,
    src_col_category: str,
    src_col_particulars: str,
    src_col_value: str,
    tgt_file: str,
    sheet_match_mode: str = "cat_in_sheet",
    tgt_col_lookup: str = "",
    tgt_col_write: str = "",
    tgt_header_row: int = 1,
    case_sensitive: bool = False,
) -> tuple[bool, str]:
    """
    Push values from a master sheet into a multi-sheet target workbook.

    Parameters
    ----------
    src_file_path      : path to master workbook
    src_sheet          : sheet in master (empty = first sheet)
    src_col_category   : header/letter/index of the category column in master
    src_col_particulars: header/letter/index of the particulars/lookup column in master
    src_col_value      : header/letter/index of the value column in master
    tgt_file           : path to target workbook (multiple sheets)
    sheet_match_mode   : how to match category → sheet name
    tgt_col_lookup     : column in target sheets to search for the particulars value
    tgt_col_write      : column in target sheets to write the value into
    tgt_header_row     : 1-based header row index in target sheets (default 1)
    case_sensitive     : whether the particulars lookup is case-sensitive
    """

    # ── Validation ────────────────────────────────────────────────────────────
    if not src_file_path or not os.path.exists(src_file_path):
        return False, f"Master file not found: {src_file_path}"
    if not tgt_file or not os.path.exists(tgt_file):
        return False, f"Target file not found: {tgt_file}"
    if not src_col_category:
        return False, "Category column (master) is required"
    if not src_col_particulars:
        return False, "Particulars column (master) is required"
    if not src_col_value:
        return False, "Value column (master) is required"
    if not tgt_col_lookup:
        return False, "Lookup column (target) is required"
    if not tgt_col_write:
        return False, "Write column (target) is required"

    # ── Read master ───────────────────────────────────────────────────────────
    try:
        df = wb_cache.read_excel(
            src_file_path,
            sheet_name=src_sheet or 0,
            dtype=str,
        )
    except Exception as exc:
        return False, f"Could not read master sheet: {exc}"

    df.columns = [str(c).strip() for c in df.columns]

    cat_col  = _find_df_col(df, src_col_category)
    part_col = _find_df_col(df, src_col_particulars)
    val_col  = _find_df_col(df, src_col_value)

    missing = []
    if not cat_col:  missing.append(f"category '{src_col_category}'")
    if not part_col: missing.append(f"particulars '{src_col_particulars}'")
    if not val_col:  missing.append(f"value '{src_col_value}'")
    if missing:
        return False, (
            f"Column(s) not found in master: {', '.join(missing)}. "
            f"Available columns: {list(df.columns)}"
        )

    # Drop rows where both category and particulars are empty
    df = df.dropna(subset=[cat_col, part_col], how="all")

    # ── Load target workbook ──────────────────────────────────────────────────
    try:
        tgt_wb = wb_cache.load(tgt_file)
    except Exception as exc:
        return False, f"Could not open target file: {exc}"

    sheet_names = tgt_wb.sheetnames

    updated        = 0
    skipped        = 0
    not_found_msgs: list[str] = []

    for _, row in df.iterrows():
        category   = str(row[cat_col]).strip()  if pd.notna(row[cat_col])  else ""
        particular = str(row[part_col]).strip() if pd.notna(row[part_col]) else ""
        value      = row[val_col]               if pd.notna(row[val_col])  else ""

        if not category or not particular:
            skipped += 1
            continue

        # 1. Match sheet
        matched = _match_sheet(category, sheet_names, sheet_match_mode)
        if not matched:
            not_found_msgs.append(f"No sheet matched category '{category}'")
            skipped += 1
            continue

        ws = tgt_wb[matched]

        # 2. Resolve lookup / write column indices in this sheet
        lookup_idx = _resolve_col_idx(ws, tgt_col_lookup)
        write_idx  = _resolve_col_idx(ws, tgt_col_write)

        if not lookup_idx:
            not_found_msgs.append(
                f"Lookup col '{tgt_col_lookup}' not found in sheet '{matched}'"
            )
            skipped += 1
            continue
        if not write_idx:
            not_found_msgs.append(
                f"Write col '{tgt_col_write}' not found in sheet '{matched}'"
            )
            skipped += 1
            continue

        # 3. Search lookup column for the particulars value
        part_cmp = particular if case_sensitive else particular.lower()
        found_row = None
        for tgt_row in ws.iter_rows(min_row=tgt_header_row + 1):
            cell_val = tgt_row[lookup_idx - 1].value
            if cell_val is None:
                continue
            cv = str(cell_val).strip()
            cv_cmp = cv if case_sensitive else cv.lower()
            if cv_cmp == part_cmp:
                found_row = tgt_row
                break

        if found_row is None:
            not_found_msgs.append(
                f"'{particular}' not found in '{matched}' (col {tgt_col_lookup})"
            )
            skipped += 1
            continue

        # 4. Write value
        found_row[write_idx - 1].value = value
        updated += 1

    # ── Save back ─────────────────────────────────────────────────────────────
    wb_cache.save(tgt_wb, tgt_file)

    summary = f"Updated {updated} cell(s)"
    if skipped:
        summary += f", skipped {skipped}"
    if not_found_msgs:
        preview = not_found_msgs[:5]
        if len(not_found_msgs) > 5:
            preview.append(f"…and {len(not_found_msgs) - 5} more")
        summary += f". Unmatched: {'; '.join(preview)}"

    return True, summary
