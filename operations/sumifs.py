"""
SUMIFS operations — pandas groupby implementation.
Reads source sheet, groups by lookup key, sums a column, then writes
the result back into the target sheet. Fast even on 50K+ rows.
NO win32com. NO Excel formula strings.
"""

import pandas as pd
from openpyxl.utils import column_index_from_string, get_column_letter
from core import wb_cache


def _col_to_idx(col: str) -> int:
    """Accept column letter ('B') or 1-based integer string ('2'). Returns 1-based int."""
    col = str(col).strip()
    if col.isdigit():
        return int(col)
    return column_index_from_string(col.upper())


def run_sumifs_step(
    target_file_path: str,
    target_sheet: str,
    source_file_path: str,
    source_sheet: str,
    src_lookup_col: str,        # col in source to match on   (e.g. "D" or "4")
    tgt_lookup_col: str,        # col in target to match from (e.g. "K" or "11")
    sum_col: str,               # col in source to sum        (e.g. "B" or "2")
    output_col: str,            # col in target to write into (e.g. "L")
    src_header_row: int = 1,
    tgt_header_row: int = 1,
    # Legacy params kept for backward compat (ignored)
    source_workbook_label=None,
    source_sheet_name=None,
    sum_range=None,
    criteria=None,
    start_row=None,
    end_row=None,
    convert_to_values=False,
) -> tuple[bool, str]:
    """
    Group-sum source data and write into the target sheet.

    Algorithm:
      1. Read source sheet → group by src_lookup_col → sum sum_col
         → produces a dict  {lookup_key: total}
      2. Open target workbook with openpyxl (preserves formatting / other columns)
      3. For every data row in target: look up tgt_lookup_col value in the dict
         → write result to output_col. Unmatched rows get 0.
    """
    try:
        # ── 1. Build the grouped-sum lookup dict from source ──────────────────
        df_src = wb_cache.read_excel(
            source_file_path,
            sheet_name=source_sheet,
            header=src_header_row - 1,
        )
        df_src.columns = [str(c).strip() for c in df_src.columns]

        src_idx = _col_to_idx(src_lookup_col)   # 1-based
        sum_idx = _col_to_idx(sum_col)           # 1-based

        # Access by positional index (header-agnostic — works even with messy headers)
        src_key_series = df_src.iloc[:, src_idx - 1].astype(str).str.strip()
        sum_series     = pd.to_numeric(df_src.iloc[:, sum_idx - 1], errors="coerce").fillna(0)

        grouped = (
            pd.DataFrame({"key": src_key_series, "val": sum_series})
            .groupby("key", sort=False)["val"]
            .sum()
            .to_dict()
        )

        # ── 2. Open target workbook with openpyxl ─────────────────────────────
        wb  = wb_cache.load(target_file_path)
        ws  = wb[target_sheet] if target_sheet and target_sheet in wb.sheetnames else wb.active

        tgt_idx = _col_to_idx(tgt_lookup_col)   # 1-based
        out_idx = _col_to_idx(output_col)        # 1-based

        first_data_row = int(tgt_header_row) + 1
        last_row       = ws.max_row

        # ── 3. Write results row by row ────────────────────────────────────────
        written = matched = 0
        for row in range(first_data_row, last_row + 1):
            raw_key = ws.cell(row=row, column=tgt_idx).value
            if raw_key is None:
                ws.cell(row=row, column=out_idx).value = 0
                written += 1
                continue
            key   = str(raw_key).strip()
            total = grouped.get(key, None)
            ws.cell(row=row, column=out_idx).value = total if total is not None else 0
            written += 1
            if total is not None:
                matched += 1

        wb_cache.save(wb, target_file_path)
        wb.close()

        out_letter = get_column_letter(out_idx)
        return True, (
            f"SUMIFS complete → {out_letter}{first_data_row}:{out_letter}{last_row} | "
            f"{matched}/{written} rows matched | "
            f"{len(grouped)} unique keys in source"
        )

    except Exception as e:
        import traceback
        return False, f"Error in SUMIFS: {str(e)}\n{traceback.format_exc()}"
