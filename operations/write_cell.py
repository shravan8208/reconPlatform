"""
Write to Cell(s) operation.
Write a value to a specific cell or fill an entire column with the same value.
Pure Python/openpyxl implementation - NO win32com!

Supports two scopes:
  single       — write into the selected sheet only (classic behaviour)
  all_sheets   — write the same cell/column into every sheet in the workbook,
                 with optional include / exclude sheet filters
"""

import openpyxl
from openpyxl.utils import column_index_from_string
from core import wb_cache


def _apply_write_to_sheet(ws, mode: str, value, cell_ref: str,
                          column: str, start_row: int, end_row) -> tuple:
    """
    Apply a single-cell or column write to one worksheet.
    Returns (cells_written: int, detail_msg: str).
    """
    if mode == "single_cell":
        if not cell_ref:
            return 0, "skipped (no cell ref)"
        old_value = ws[cell_ref].value
        ws[cell_ref].value = value
        return 1, f"{cell_ref}: '{old_value}' → '{value}'"

    elif mode == "column":
        if not column:
            return 0, "skipped (no column)"
        col_idx = column_index_from_string(column.strip().upper())
        start_row_num = int(start_row)
        end_row_num = ws.max_row if (end_row == "last" or end_row == "") else int(end_row)
        count = 0
        for r in range(start_row_num, end_row_num + 1):
            ws.cell(row=r, column=col_idx).value = value
            count += 1
        return count, f"col {column} rows {start_row_num}-{end_row_num}: {count} cell(s)"

    return 0, f"unknown mode '{mode}'"


def run_write_cell_step(
    file_path: str,
    sheet_name: str,
    mode: str,            # "single_cell" or "column"
    value: str,
    # For single cell mode:
    cell_ref: str = "",   # e.g., "A5", "B10"
    # For column mode:
    column: str = "",     # Column letter, e.g., "A", "B"
    start_row: int = 2,
    end_row = "last",     # "last" or a number
    header_row: int = 1,
    # All-sheets scope:
    scope: str = "single",          # "single" | "all_sheets"
    include_sheets: list = None,    # only process these sheets
    exclude_sheets: list = None,    # skip these sheets
):
    """
    Write a value to a specific cell / column — in one sheet or across all sheets.

    scope = "single"     → classic behaviour, writes to sheet_name only.
    scope = "all_sheets" → writes to every sheet in the workbook.
                           include_sheets / exclude_sheets restrict which sheets.
    """
    try:
        wb = wb_cache.load(file_path)

        # ── Resolve target sheets ──────────────────────────────────────────────
        if scope == "all_sheets":
            inc = [s.strip() for s in (include_sheets or []) if s.strip()]
            exc = {s.strip().lower() for s in (exclude_sheets or []) if s.strip()}
            if inc:
                target_sheets = [s for s in wb.sheetnames if s in inc]
            else:
                target_sheets = [s for s in wb.sheetnames if s.strip().lower() not in exc]
        else:
            if sheet_name and sheet_name in wb.sheetnames:
                target_sheets = [sheet_name]
            else:
                target_sheets = [wb.active.title]

        if not target_sheets:
            return False, "No sheets to process."

        # ── Write ──────────────────────────────────────────────────────────────
        total_cells = 0
        details: list[str] = []
        errors:  list[str] = []

        for sn in target_sheets:
            ws = wb[sn]
            try:
                n, msg = _apply_write_to_sheet(
                    ws, mode, value, cell_ref, column, start_row, end_row
                )
                total_cells += n
                prefix = f"[{sn}] " if scope == "all_sheets" else ""
                details.append(f"{prefix}{msg}")
            except Exception as exc:
                errors.append(f"[{sn}] {exc}")

        wb_cache.save(wb, file_path)
        wb.close()

        if errors and not details:
            return False, " | ".join(errors)

        sheet_label = (
            f"{len(target_sheets)} sheet(s)" if scope == "all_sheets" else sheet_name
        )
        summary = f"Wrote '{value}' to {total_cells} cell(s) across {sheet_label}."
        if scope == "all_sheets" and details:
            summary += "\n" + "\n".join(details)
        elif details:
            summary = details[0]  # single-sheet: keep concise original message
        if errors:
            summary += "\nWarnings: " + " | ".join(errors)

        return True, summary

    except Exception as e:
        return False, f"Error writing to cell(s): {str(e)}"

