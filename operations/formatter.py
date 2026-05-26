"""
Formatting & Beautification operations.
Applies static and conditional cell formatting to Excel sheets.
Pure openpyxl — no win32com.

Static mode  : apply a style to fixed rows / columns / ranges (headers, templates).
Conditional  : evaluate a condition per row and apply style to matching cells/rows.
"""

import re
import pandas as pd
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.styles.colors import Color as _OXLColor
from openpyxl.utils import column_index_from_string
from core import wb_cache


# ---------------------------------------------------------------------------
# Catalogues
# ---------------------------------------------------------------------------

BORDER_STYLE_OPTS = {
    "none":   None,
    "thin":   "thin",
    "medium": "medium",
    "thick":  "thick",
    "double": "double",
    "dashed": "dashed",
    "dotted": "dotted",
}

BORDER_SIDE_OPTS = ["all", "outer", "top", "bottom", "left", "right", "top_bottom"]

ALIGN_OPTS = {
    "":        None,
    "left":    "left",
    "center":  "center",
    "right":   "right",
    "justify": "justify",
    "general": "general",
}

COND_OPERATORS = {
    "gt":           "> greater than",
    "lt":           "< less than",
    "gte":          ">= greater or equal",
    "lte":          "<= less or equal",
    "eq":           "= equals",
    "ne":           "!= not equals",
    "contains":     "contains",
    "not_contains": "does not contain",
    "starts_with":  "starts with",
    "ends_with":    "ends with",
    "is_empty":     "is empty / blank",
    "is_not_empty": "is not empty / blank",
    "between":      "between (value and value2)",
}

COND_NO_VALUE   = {"is_empty", "is_not_empty"}
COND_TWO_VALUE  = {"between"}
COND_NUMERIC    = {"gt", "lt", "gte", "lte", "between"}

APPLY_TO_OPTS = {
    "row":  "Entire row",
    "cell": "Matched cell only",
    "cols": "Specific columns",
}


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _hex_to_argb(hex_color: str) -> str:
    """Convert #RRGGBB / RRGGBB / #RGB to FFRRGGBB.  Returns '' on failure."""
    h = str(hex_color or "").strip().lstrip("#").upper()
    if len(h) == 3:
        h = h[0] * 2 + h[1] * 2 + h[2] * 2
    if len(h) == 6:
        return "FF" + h
    if len(h) == 8:
        return h
    return ""


def _apply_style_to_cell(cell, style: dict):
    """Apply a complete style dict to a single openpyxl cell.  Merges with existing style."""
    if not style:
        return

    # ── Font ────────────────────────────────────────────────────────────────
    bold       = style.get("bold")        # True / False / None
    italic     = style.get("italic")
    underline  = style.get("underline")
    strike     = style.get("strike")
    font_size  = style.get("font_size")   # int or None
    font_color = style.get("font_color", "")  # hex string

    has_font = any(x is not None for x in [bold, italic, underline, strike, font_size]) or font_color
    if has_font:
        ex = cell.font or Font()
        fc_argb = _hex_to_argb(font_color) if font_color else None
        cell.font = Font(
            name      = ex.name,
            bold      = bool(bold)     if bold      is not None else ex.bold,
            italic    = bool(italic)   if italic    is not None else ex.italic,
            underline = "single"       if underline else (ex.underline if underline is None else None),
            strike    = bool(strike)   if strike    is not None else ex.strike,
            size      = int(font_size) if font_size else ex.size,
            color     = _OXLColor(rgb=fc_argb) if fc_argb else ex.color,
        )

    # ── Fill ────────────────────────────────────────────────────────────────
    bg_color = style.get("bg_color", "")
    if bg_color:
        argb = _hex_to_argb(bg_color)
        if argb:
            cell.fill = PatternFill(fill_type="solid", fgColor=argb)

    # ── Border ──────────────────────────────────────────────────────────────
    border_style_key = str(style.get("border_style", "none")).lower()
    border_side_key  = str(style.get("border_sides",  "all")).lower()
    border_opxl_style = BORDER_STYLE_OPTS.get(border_style_key, None)

    if border_opxl_style:
        active = Side(style=border_opxl_style)
        none_s = Side(style=None)
        ex_b   = cell.border or Border()

        sides_active: set = set()
        if border_side_key in ("all", "outer"):
            sides_active = {"top", "bottom", "left", "right"}
        elif border_side_key == "top_bottom":
            sides_active = {"top", "bottom"}
        else:
            sides_active = {border_side_key}  # single side

        cell.border = Border(
            top    = active if "top"    in sides_active else ex_b.top,
            bottom = active if "bottom" in sides_active else ex_b.bottom,
            left   = active if "left"   in sides_active else ex_b.left,
            right  = active if "right"  in sides_active else ex_b.right,
        )

    # ── Alignment ───────────────────────────────────────────────────────────
    h_align   = ALIGN_OPTS.get(str(style.get("h_align", "")).lower(), None)
    wrap_text = style.get("wrap_text")  # True / False / None

    if h_align or wrap_text is not None:
        ex_a = cell.alignment or Alignment()
        cell.alignment = Alignment(
            horizontal = h_align   if h_align   else ex_a.horizontal,
            wrap_text  = bool(wrap_text) if wrap_text is not None else ex_a.wrap_text,
        )


def _resolve_col(df: pd.DataFrame, col_spec: str, col_mode: str):
    """
    Resolve col_spec to (col_name, 0-based-index) in df.
    Returns (None, None) on failure.
    """
    s = str(col_spec or "").strip()
    if not s:
        return None, None

    cols = list(df.columns)

    if col_mode == "header":
        if s in cols:
            return s, cols.index(s)
        lm = {str(c).lower(): c for c in cols}
        if s.lower() in lm:
            c = lm[s.lower()]
            return c, cols.index(c)
        return None, None

    # letter / number mode
    if s.isdigit():
        idx = int(s) - 1
    else:
        try:
            idx = column_index_from_string(s.upper()) - 1
        except Exception:
            return None, None

    if 0 <= idx < len(cols):
        return cols[idx], idx
    return None, None


def _build_mask(series: pd.Series, op: str, value: str, value2: str = "") -> pd.Series:
    """Return a boolean Series for one condition."""
    s_str = series.astype(str).str.strip()
    val   = str(value or "")

    if op == "is_empty":
        return series.isna() | (s_str == "") | (s_str.isin(["None", "nan", "NaN"]))
    if op == "is_not_empty":
        return ~(series.isna() | (s_str == "") | (s_str.isin(["None", "nan", "NaN"])))

    if op in COND_NUMERIC:
        try:
            nv = float(val)
            ns = pd.to_numeric(series, errors="coerce")
            if op == "gt":  return ns > nv
            if op == "lt":  return ns < nv
            if op == "gte": return ns >= nv
            if op == "lte": return ns <= nv
            if op == "between":
                try:
                    nv2 = float(str(value2 or ""))
                    return (ns >= nv) & (ns <= nv2)
                except ValueError:
                    return pd.Series([False] * len(series), index=series.index)
        except (ValueError, TypeError):
            return pd.Series([False] * len(series), index=series.index)

    sl = s_str.str.lower()
    vl = val.lower()
    if op == "eq":           return sl == vl
    if op == "ne":           return sl != vl
    if op == "contains":     return sl.str.contains(vl, na=False, regex=False)
    if op == "not_contains": return ~sl.str.contains(vl, na=False, regex=False)
    if op == "starts_with":  return sl.str.startswith(vl, na=False)
    if op == "ends_with":    return sl.str.endswith(vl, na=False)

    return sl == vl  # fallback = equals


# ---------------------------------------------------------------------------
# Static formatting
# ---------------------------------------------------------------------------

def _apply_static_rule(ws, target_type: str, target_spec: str, style: dict) -> int:
    """
    Apply style to a fixed target in ws.
    target_type: "row" | "col" | "range"
    target_spec: "1,2,3" | "A,B" | "A1:Z5"
    Returns number of cells styled.
    """
    styled = 0

    if target_type == "range":
        try:
            for row in ws[target_spec.strip()]:
                for cell in (row if hasattr(row, "__iter__") else [row]):
                    _apply_style_to_cell(cell, style)
                    styled += 1
        except Exception:
            pass

    elif target_type == "row":
        for part in target_spec.split(","):
            part = part.strip()
            if not part:
                continue
            try:
                row_idx = int(part)
                for cell in ws[row_idx]:
                    _apply_style_to_cell(cell, style)
                    styled += 1
            except Exception:
                pass

    elif target_type == "col":
        for part in target_spec.split(","):
            part = part.strip().upper()
            if not part:
                continue
            try:
                col_idx = int(part) if part.isdigit() else column_index_from_string(part)
                for row_cells in ws.iter_rows(min_col=col_idx, max_col=col_idx,
                                               min_row=1, max_row=ws.max_row):
                    for cell in row_cells:
                        _apply_style_to_cell(cell, style)
                        styled += 1
            except Exception:
                pass

    return styled


# ---------------------------------------------------------------------------
# Conditional formatting (manual row-by-row via pandas)
# ---------------------------------------------------------------------------

def _apply_conditional_rules(ws, df: pd.DataFrame, rules: list,
                              col_mode: str, header_row: int) -> list:
    """
    Apply conditional format rules to ws.
    Returns list of (rule_idx, rows_matched, summary_str).
    """
    results = []

    for r_idx, rule in enumerate(rules, start=1):
        check_col_spec   = str(rule.get("check_col",   "") or "").strip()
        op               = rule.get("operator", "gt")
        val              = str(rule.get("value",  "") or "")
        val2             = str(rule.get("value2", "") or "")
        apply_to         = str(rule.get("apply_to", "row")).lower()
        target_cols_spec = str(rule.get("target_cols", "") or "")
        style            = rule.get("style") or {}

        if not check_col_spec:
            results.append((r_idx, 0, f"Rule {r_idx}: no check column — skipped"))
            continue

        col_name, col_idx0 = _resolve_col(df, check_col_spec, col_mode)
        if col_name is None:
            results.append((r_idx, 0, f"Rule {r_idx}: column '{check_col_spec}' not found — skipped"))
            continue

        mask = _build_mask(df[col_name], op, val, val2)
        matched_indices = df.index[mask].tolist()
        n_matched = len(matched_indices)

        # Resolve target columns for "cols" mode (1-based xl column numbers)
        target_xl_cols = []
        if apply_to == "cols" and target_cols_spec:
            for part in target_cols_spec.split(","):
                part = part.strip()
                if not part:
                    continue
                _, c_idx0 = _resolve_col(df, part, col_mode)
                if c_idx0 is not None:
                    target_xl_cols.append(c_idx0 + 1)  # 1-based

        for df_row_idx in matched_indices:
            # pandas index 0 → header_row + 1 in Excel (1-based)
            xl_row = df_row_idx + header_row + 1

            if apply_to == "row":
                for cell in ws[xl_row]:
                    _apply_style_to_cell(cell, style)

            elif apply_to == "cell":
                cell = ws.cell(row=xl_row, column=col_idx0 + 1)
                _apply_style_to_cell(cell, style)

            elif apply_to == "cols" and target_xl_cols:
                for xl_col in target_xl_cols:
                    cell = ws.cell(row=xl_row, column=xl_col)
                    _apply_style_to_cell(cell, style)

            else:  # default → whole row
                for cell in ws[xl_row]:
                    _apply_style_to_cell(cell, style)

        op_label = COND_OPERATORS.get(op, op)
        val_str  = "" if op in COND_NO_VALUE else f" '{val}'"
        if op in COND_TWO_VALUE:
            val_str = f" '{val}' – '{val2}'"
        summary = (
            f"Rule {r_idx}: [{check_col_spec}] {op_label}{val_str}"
            f"  →  {n_matched} row(s) formatted"
        )
        results.append((r_idx, n_matched, summary))

    return results


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

def run_format_step(
    file_path: str,
    sheet_name: str,
    format_mode: str  = "static",   # "static" | "conditional"
    header_row: int   = 1,
    col_mode: str     = "letter",   # "letter" | "header"
    static_rules: list  = None,     # [{target_type, target_spec, style}]
    cond_rules:   list  = None,     # [{check_col, operator, value, value2, apply_to, target_cols, style}]
) -> tuple:
    """
    Apply static or conditional formatting to an Excel sheet.
    Returns (success: bool, message: str).
    """
    try:
        wb = wb_cache.load(file_path)
        if sheet_name not in wb.sheetnames:
            return False, f"Sheet '{sheet_name}' not found in {file_path}"
        ws = wb[sheet_name]

        summaries = []
        total_affected = 0

        if format_mode == "static":
            for r_idx, rule in enumerate(static_rules or [], start=1):
                t_type = str(rule.get("target_type", "row")).lower()
                t_spec = str(rule.get("target_spec", "") or "")
                sty    = rule.get("style") or {}
                if not t_spec:
                    continue
                n = _apply_static_rule(ws, t_type, t_spec, sty)
                total_affected += n
                summaries.append(f"Rule {r_idx}: {t_type.upper()} {t_spec}  →  {n} cell(s)")

        else:  # conditional
            df = wb_cache.read_excel(file_path, sheet_name=sheet_name, header=header_row - 1)
            df.columns = [str(c).strip() for c in df.columns]
            rule_results = _apply_conditional_rules(ws, df, cond_rules or [], col_mode, header_row)
            for _, n, s in rule_results:
                total_affected += n
                summaries.append(s)

        wb_cache.save(wb, file_path)
        wb.close()

        mode_label = "Static" if format_mode == "static" else "Conditional"
        msg = (
            f"{mode_label} Format: {len(summaries)} rule(s), "
            f"{total_affected} cell(s)/row(s) affected.\n"
            + "\n".join(summaries)
        )
        return True, msg

    except Exception as e:
        return False, f"Format step error: {e}"
