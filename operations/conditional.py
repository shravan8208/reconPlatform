"""
Conditional write operations.
Pure Python/pandas implementation - NO win32com!

Supports:
  - 13 comparison operators: equals, not_equals, contains, not_contains,
    starts_with, ends_with, greater_than, less_than, greater_than_eq,
    less_than_eq, is_empty, is_not_empty, regex
  - Columns identified by letter/number OR by header name
"""

import re
import pandas as pd
import openpyxl
from openpyxl.utils import column_index_from_string
from openpyxl.utils.dataframe import dataframe_to_rows
from core import wb_cache


# ── Operator registry ────────────────────────────────────────────────────────

OPERATORS = {
    "equals":          "Equals (exact)",
    "not_equals":      "Not Equals",
    "contains":        "Contains",
    "not_contains":    "Does Not Contain",
    "starts_with":     "Starts With",
    "ends_with":       "Ends With",
    "greater_than":    "Greater Than  ( > )",
    "less_than":       "Less Than  ( < )",
    "greater_than_eq": "Greater Than or Equal  ( >= )",
    "less_than_eq":    "Less Than or Equal  ( <= )",
    "is_empty":        "Is Empty / Blank",
    "is_not_empty":    "Is Not Empty / Blank",
    "regex":           "Matches Regex",
}

# Operators that do NOT need a comparison value
NO_VALUE_OPERATORS = {"is_empty", "is_not_empty"}

# Operators that require numeric comparison
NUMERIC_OPERATORS = {"greater_than", "less_than", "greater_than_eq", "less_than_eq"}


# ── Helpers ──────────────────────────────────────────────────────────────────

def _resolve_column(df: pd.DataFrame, col_spec: str, col_mode: str, header_row: int):
    """
    Resolve a column spec to a DataFrame column name.

    col_mode = "letter"  — col_spec is a column letter (A, B, …) or 1-based integer
    col_mode = "header"  — col_spec is the exact header string

    Returns the resolved column name (as it appears in df.columns), or raises
    a ValueError if not found.
    """
    if col_mode == "header":
        name = str(col_spec).strip()
        if name in df.columns:
            return name
        # case-insensitive fallback
        lower_map = {str(c).strip().lower(): c for c in df.columns}
        if name.lower() in lower_map:
            return lower_map[name.lower()]
        raise ValueError(f"Header '{name}' not found in sheet columns: {list(df.columns)}")
    else:
        # letter or number mode
        s = str(col_spec).strip()
        if not s:
            raise ValueError("Column specification is empty")
        try:
            if s.isdigit():
                idx = int(s) - 1          # 1-based → 0-based
            else:
                idx = column_index_from_string(s.upper()) - 1
        except Exception:
            raise ValueError(f"Cannot parse column spec '{s}' as letter or number")
        cols = list(df.columns)
        if idx < 0 or idx >= len(cols):
            raise ValueError(f"Column index {idx + 1} out of range (sheet has {len(cols)} columns)")
        return cols[idx]


def _build_mask(series: pd.Series, operator: str, condition_value: str) -> pd.Series:
    """
    Build a boolean mask applying the chosen operator against condition_value.
    String comparisons are case-insensitive.
    Numeric comparisons coerce the series and the value to float.
    """
    op = operator or "equals"
    val = str(condition_value) if condition_value is not None else ""
    s_str = series.astype(str).str.strip()

    if op == "equals":
        return s_str.str.lower() == val.lower()

    if op == "not_equals":
        return s_str.str.lower() != val.lower()

    if op == "contains":
        return s_str.str.contains(val, case=False, na=False, regex=False)

    if op == "not_contains":
        return ~s_str.str.contains(val, case=False, na=False, regex=False)

    if op == "starts_with":
        return s_str.str.lower().str.startswith(val.lower())

    if op == "ends_with":
        return s_str.str.lower().str.endswith(val.lower())

    if op == "is_empty":
        return series.isna() | (s_str == "") | (s_str == "None") | (s_str == "nan")

    if op == "is_not_empty":
        return ~(series.isna() | (s_str == "") | (s_str == "None") | (s_str == "nan"))

    if op == "regex":
        try:
            return series.astype(str).str.contains(val, case=False, na=False, regex=True)
        except re.error as e:
            raise ValueError(f"Invalid regex pattern '{val}': {e}")

    if op in NUMERIC_OPERATORS:
        try:
            num_val = float(val)
        except (ValueError, TypeError):
            raise ValueError(
                f"Operator '{op}' requires a numeric value; got '{val}'"
            )
        numeric_series = pd.to_numeric(series, errors="coerce")
        if op == "greater_than":
            return numeric_series > num_val
        if op == "less_than":
            return numeric_series < num_val
        if op == "greater_than_eq":
            return numeric_series >= num_val
        if op == "less_than_eq":
            return numeric_series <= num_val

    raise ValueError(f"Unknown operator '{op}'")


# ── Internal per-rule applicator ─────────────────────────────────────────────

def _apply_rule(df: pd.DataFrame, rule: dict, col_mode: str, header_row: int):
    """
    Apply one rule dict to df in-place.
    Returns (rows_affected: int, summary_str: str).
    Raises ValueError on bad config.
    """
    cond_col    = str(rule.get("cond_col", "")).strip()
    operator    = rule.get("operator", "equals")
    cond_val    = str(rule.get("cond_val", "") or "")
    target_col  = str(rule.get("target_col", "")).strip()
    write_value = rule.get("write_val", "")

    if not cond_col or not target_col:
        raise ValueError("Each rule must have a condition column and a target column.")

    cond_col_name = _resolve_column(df, cond_col, col_mode, header_row)

    try:
        tgt_col_name = _resolve_column(df, target_col, col_mode, header_row)
    except ValueError:
        if col_mode == "letter":
            df[str(target_col).strip()] = ""
            tgt_col_name = str(target_col).strip()
        else:
            raise

    mask = _build_mask(df[cond_col_name], operator, cond_val)
    rows_affected = int(mask.sum())
    if rows_affected:
        df.loc[mask, tgt_col_name] = write_value

    op_label = OPERATORS.get(operator, operator)
    val_str  = "" if operator in NO_VALUE_OPERATORS else f" '{cond_val}'"
    summary  = (
        f"[{cond_col_name}] {op_label}{val_str} → "
        f"[{tgt_col_name}] = '{write_value}'  ({rows_affected} row(s))"
    )
    return rows_affected, summary


# ── Main function ────────────────────────────────────────────────────────────

def run_conditional_write_step(
    file_path,
    sheet_name,
    condition_col       = None,   # single-rule legacy param
    condition_value     = None,   # single-rule legacy param
    target_col          = None,   # single-rule legacy param
    write_value         = None,   # single-rule legacy param
    operator: str       = "equals",
    header_row: int     = 1,
    col_mode: str       = "letter",
    rules: list         = None,   # multi-rule list; overrides single-rule params when provided
):
    """
    Apply one or more conditional-write rules to a sheet in a single pass.

    Multi-rule mode (preferred):
        Pass `rules` as a list of dicts, each with keys:
            cond_col, operator, cond_val, target_col, write_val

    Single-rule mode (backward-compatible):
        Pass condition_col, condition_value, target_col, write_value, operator.

    Both modes share: file_path, sheet_name, header_row, col_mode.
    """
    try:
        # ── Normalise rules list ──────────────────────────────────────────────
        if rules:
            rule_list = [r for r in rules if r]
        else:
            # Legacy single-rule → wrap in list
            rule_list = [{
                "cond_col":   condition_col   or "",
                "operator":   operator,
                "cond_val":   condition_value or "",
                "target_col": target_col      or "",
                "write_val":  write_value     or "",
            }]

        if not rule_list:
            return False, "No rules defined."

        # ── Load DataFrame ────────────────────────────────────────────────────
        df = wb_cache.read_excel(
            file_path,
            sheet_name=sheet_name,
            header=header_row - 1,
        )
        df.columns = [str(c).strip() for c in df.columns]

        # ── Apply all rules sequentially ──────────────────────────────────────
        total_affected = 0
        summaries      = []
        rule_errors    = []

        for idx, rule in enumerate(rule_list, start=1):
            try:
                n, summary = _apply_rule(df, rule, col_mode, header_row)
                total_affected += n
                summaries.append(f"Rule {idx}: {summary}")
            except Exception as re:
                rule_errors.append(f"Rule {idx} error: {re}")

        if rule_errors and not summaries:
            return False, " | ".join(rule_errors)

        # ── Write back ────────────────────────────────────────────────────────
        wb = wb_cache.load(file_path)
        if sheet_name in wb.sheetnames:
            ws_pos = wb.sheetnames.index(sheet_name)
            del wb[sheet_name]
            ws = wb.create_sheet(sheet_name, ws_pos)
        else:
            ws = wb.create_sheet(sheet_name)

        for r_idx, row in enumerate(
            dataframe_to_rows(df, index=False, header=True), start=1
        ):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)

        wb_cache.save(wb, file_path)
        wb.close()

        msg = (
            f"Conditional write: {len(rule_list)} rule(s), "
            f"{total_affected} total row(s) updated.\n"
            + "\n".join(summaries)
        )
        if rule_errors:
            msg += "\nWarnings: " + " | ".join(rule_errors)

        return True, msg

    except Exception as e:
        return False, f"Conditional write error: {e}"
