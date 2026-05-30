"""
Data manipulation operations: forward fill, insert, delete rows/columns.
Pure openpyxl implementation - NO win32com!
"""

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
from typing import List, Optional, Tuple, Union
from core import wb_cache


def run_forward_fill_step(file_path, sheet_name, col_letter):
    """
    Forward fill (fill down) operation on a column.
    
    Args:
        file_path: Path to Excel file
        sheet_name: Sheet name
        col_letter: Column letter to fill
    
    Returns:
        Tuple of (success: bool, message: str)
    """
    try:
        wb = wb_cache.load(file_path)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        
        col_idx = column_index_from_string(col_letter)
        max_row = ws.max_row
        
        last_value = None
        filled_count = 0
        
        for row in range(1, max_row + 1):
            cell = ws.cell(row=row, column=col_idx)
            
            if cell.value is not None and cell.value != '':
                last_value = cell.value
            elif last_value is not None:
                cell.value = last_value
                filled_count += 1
        
        wb_cache.save(wb, file_path)
        wb.close()
        
        return True, f"Forward filled column {col_letter}, {filled_count} cells filled"
        
    except Exception as e:
        return False, f"Error in forward fill: {str(e)}"


def run_insert_delete_step(file_path, sheet_name, op_type, axis, start_idx, end_idx_str):
    """
    Insert or delete rows/columns.
    
    Args:
        file_path: Path to Excel file
        sheet_name: Sheet name
        op_type: 'insert' or 'delete'
        axis: 'Row' or 'Column'
        start_idx: Starting index (1-based)
        end_idx_str: Ending index or 'last'
    
    Returns:
        Tuple of (success: bool, message: str)
    """
    try:
        wb = wb_cache.load(file_path)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        
        start_idx = int(start_idx)
        
        if end_idx_str == 'last':
            end_idx = ws.max_row if axis.lower() == 'row' else ws.max_column
        else:
            end_idx = int(end_idx_str)
        
        count = end_idx - start_idx + 1
        
        if op_type.lower() == 'insert':
            if axis.lower() == 'row':
                ws.insert_rows(start_idx, count)
            else:
                ws.insert_cols(start_idx, count)
            action = f"Inserted {count} {axis.lower()}(s) at position {start_idx}"
        else:  # delete
            if axis.lower() == 'row':
                ws.delete_rows(start_idx, count)
            else:
                ws.delete_cols(start_idx, count)
            action = f"Deleted {count} {axis.lower()}(s) from position {start_idx}"
        
        wb_cache.save(wb, file_path)
        wb.close()
        
        return True, action
        
    except Exception as e:
        return False, f"Error in insert/delete: {str(e)}"


def _resolve_col_idx(ws, header_row: int, col_or_header: str) -> Optional[int]:
    """
    Resolve a column identifier to a 1-based column index.
    Accepts:
    - Column letter(s): "A", "AB"
    - Header name present in header_row
    """
    s = str(col_or_header or "").strip()
    if not s:
        return None

    # Try letter
    if s.isalpha() and 1 <= len(s) <= 3:
        try:
            idx = column_index_from_string(s.upper())
            if idx <= ws.max_column:
                return idx
        except Exception:
            pass

    # Try header name match (trimmed)
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=int(header_row), column=c).value
        if v is None:
            continue
        if str(v).strip() == s:
            return c
    return None


def run_delete_columns_step(
    file_path: str,
    sheet_name: str,
    columns: Union[str, List[str]],
    header_row: int = 1,
) -> Tuple[bool, str]:
    """
    Delete specific column(s) by header name or column letter.

    This supports deleting multiple non-contiguous columns. Deletions are performed
    from right-to-left to keep indices stable.
    """
    try:
        if isinstance(columns, str):
            cols = [columns]
        else:
            cols = list(columns or [])
        cols = [str(c).strip() for c in cols if str(c).strip()]
        if not cols:
            return False, "No columns selected to delete"

        wb = wb_cache.load(file_path)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

        idxs: List[int] = []
        missing: List[str] = []
        for c in cols:
            idx = _resolve_col_idx(ws, int(header_row), c)
            if idx is None:
                missing.append(c)
            else:
                if idx not in idxs:
                    idxs.append(idx)

        if missing:
            wb.close()
            return False, f"Column(s) not found to delete (row {header_row} headers): {', '.join(missing)}"

        # Delete from right to left
        for idx in sorted(idxs, reverse=True):
            ws.delete_cols(idx, 1)

        wb_cache.save(wb, file_path)
        wb.close()

        return True, f"Deleted {len(idxs)} column(s): {', '.join(cols)}"
    except Exception as e:
        return False, f"Error deleting columns: {str(e)}"


def run_clear_columns_data_step(
    file_path: str,
    sheet_name: str,
    columns: Union[str, List[str]],
    header_row: int = 1,
    end_row: Union[str, int] = "last",
) -> Tuple[bool, str]:
    """
    Clear (blank) the DATA under a column header, without deleting the column itself.

    - Keeps the header cell (at header_row) unchanged
    - Clears values from row header_row+1 through end_row (or last row)
    - Preserves formatting (only cell.value is cleared)
    """
    try:
        if isinstance(columns, str):
            cols = [columns]
        else:
            cols = list(columns or [])
        cols = [str(c).strip() for c in cols if str(c).strip()]
        if not cols:
            return False, "No columns selected to clear"

        wb = wb_cache.load(file_path)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

        idxs: List[int] = []
        missing: List[str] = []
        for c in cols:
            idx = _resolve_col_idx(ws, int(header_row), c)
            if idx is None:
                missing.append(c)
            else:
                if idx not in idxs:
                    idxs.append(idx)

        if missing:
            wb.close()
            return False, f"Column(s) not found to clear (row {header_row} headers): {', '.join(missing)}"

        first_data_row = int(header_row) + 1
        if str(end_row).lower() == "last" or end_row == "":
            last_row = ws.max_row
        else:
            last_row = int(end_row)

        if last_row < first_data_row:
            wb.close()
            return False, "End row is before the first data row"

        cleared = 0
        for r in range(first_data_row, last_row + 1):
            for cidx in idxs:
                ws.cell(row=r, column=cidx).value = None
                cleared += 1

        wb_cache.save(wb, file_path)
        wb.close()

        return True, f"Cleared data in {len(idxs)} column(s) for rows {first_data_row}:{last_row} ({cleared} cells)"
    except Exception as e:
        return False, f"Error clearing column data: {str(e)}"


def run_delete_by_condition_step(
    file_path: str,
    sheet_name: str,
    filters: list,
    filter_combine: str = "AND",
    header_row: int = 1,
) -> tuple[bool, str]:
    """
    Delete all rows in the sheet that match the given filter conditions.
    Rows are deleted from bottom to top to preserve row indices.
    """
    import pandas as pd
    from operations.filter_utils import apply_filters

    try:
        wb = wb_cache.load(file_path)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

        # Read into DataFrame to evaluate conditions
        data = list(ws.values)
        if not data or len(data) < header_row:
            return False, "Sheet is empty or header row is beyond data"

        headers = [str(c).strip() if c is not None else f"Col{i}" for i, c in enumerate(data[header_row - 1])]
        rows_data = data[header_row:]  # data rows only

        df = pd.DataFrame(rows_data, columns=headers)
        df["__ws_row__"] = list(range(header_row + 1, header_row + 1 + len(df)))  # 1-indexed worksheet rows

        df_match = apply_filters(df, filters, filter_combine)
        ws_rows_to_delete = sorted(df_match["__ws_row__"].tolist(), reverse=True)

        if not ws_rows_to_delete:
            return True, "No rows matched the conditions — nothing deleted"

        for row_idx in ws_rows_to_delete:
            ws.delete_rows(row_idx)

        wb_cache.save(wb, file_path)
        wb.close()

        return True, f"Deleted {len(ws_rows_to_delete)} row(s) matching conditions"
    except Exception as e:
        return False, f"Error in delete by condition: {str(e)}"


# ---------------------------------------------------------------------------
# Advanced Delete  (multi-sheet, delete log)
# ---------------------------------------------------------------------------

def _resolve_target_sheets(
    wb,
    sheet_scope: str,
    sheet_name: str,
    selected_sheets: list,
    include_sheets: list,
    exclude_sheets: list,
) -> list:
    """Return the ordered list of sheet names to process."""
    if sheet_scope == "selected_sheets":
        return [s for s in selected_sheets if s in wb.sheetnames]
    if sheet_scope == "all_sheets":
        inc = [s.strip() for s in (include_sheets or []) if s.strip()]
        exc = {s.strip().lower() for s in (exclude_sheets or []) if s.strip()}
        if inc:
            return [s for s in wb.sheetnames if s in inc]
        return [s for s in wb.sheetnames if s.strip().lower() not in exc]
    # default: single sheet
    return [sheet_name] if sheet_name else []


def _read_sheet_to_df(ws, header_row: int):
    """Read an openpyxl worksheet into a DataFrame (preserving ws row indices)."""
    import pandas as pd
    data = list(ws.values)
    if not data or len(data) < header_row:
        return None, []
    headers = [
        str(c).strip() if c is not None else f"Col{i}"
        for i, c in enumerate(data[header_row - 1])
    ]
    rows_data = data[header_row:]
    df = pd.DataFrame(rows_data, columns=headers)
    # Track the real worksheet row (1-indexed)
    df["__ws_row__"] = list(range(header_row + 1, header_row + 1 + len(df)))
    return df, headers


def preview_advanced_delete(
    file_path: str,
    sheet_scope: str = "single",
    sheet_name: str = "",
    selected_sheets: list = None,
    include_sheets: list = None,
    exclude_sheets: list = None,
    filters: list = None,
    filter_combine: str = "AND",
    header_row: int = 1,
) -> dict:
    """
    Return a preview dict without modifying any file.

    Returns:
        {
          "total": int,
          "by_sheet": {"SheetName": count, ...},
          "error": str | None,
        }
    """
    from operations.filter_utils import apply_filters

    result = {"total": 0, "by_sheet": {}, "error": None}
    try:
        wb = wb_cache.load(file_path)
        target_sheets = _resolve_target_sheets(
            wb, sheet_scope, sheet_name,
            selected_sheets or [], include_sheets or [], exclude_sheets or []
        )
        for sn in target_sheets:
            ws = wb[sn]
            df, _ = _read_sheet_to_df(ws, header_row)
            if df is None:
                result["by_sheet"][sn] = 0
                continue
            matched = apply_filters(df, filters or [], filter_combine)
            result["by_sheet"][sn] = len(matched)
            result["total"] += len(matched)
    except Exception as exc:
        result["error"] = str(exc)
    return result


def run_advanced_delete_step(
    file_path: str,
    sheet_scope: str = "single",
    sheet_name: str = "",
    selected_sheets: list = None,
    include_sheets: list = None,
    exclude_sheets: list = None,
    filters: list = None,
    filter_combine: str = "AND",
    header_row: int = 1,
    save_log: bool = True,
    log_path: str = "",
) -> tuple:
    """
    Delete rows matching the given filter conditions from one or more sheets.

    Parameters
    ----------
    file_path       : path to the Excel workbook
    sheet_scope     : "single" | "selected_sheets" | "all_sheets"
    sheet_name      : used when scope = "single"
    selected_sheets : used when scope = "selected_sheets"
    include_sheets  : restrict all_sheets scope to these sheet names
    exclude_sheets  : skip these sheets in all_sheets scope
    filters         : list of condition dicts (column / operator / value)
    filter_combine  : "AND" | "OR"
    header_row      : 1-based header row index
    save_log        : if True, write a delete-log Excel file
    log_path        : explicit path for the log file; auto-generated if empty
    """
    import os, datetime, pandas as pd
    from operations.filter_utils import apply_filters

    if not filters:
        return False, "No conditions defined — nothing to delete."

    try:
        wb = wb_cache.load(file_path)
        target_sheets = _resolve_target_sheets(
            wb, sheet_scope, sheet_name,
            selected_sheets or [], include_sheets or [], exclude_sheets or []
        )
        if not target_sheets:
            return False, "No sheets to process."

        log_rows: list[dict] = []
        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        cond_label = _describe_filters(filters, filter_combine)

        total_deleted = 0
        sheet_counts: dict[str, int] = {}

        for sn in target_sheets:
            ws = wb[sn]
            df, headers = _read_sheet_to_df(ws, header_row)
            if df is None:
                sheet_counts[sn] = 0
                continue

            matched = apply_filters(df, filters or [], filter_combine)
            ws_rows = sorted(matched["__ws_row__"].tolist(), reverse=True)

            # Capture log data BEFORE deleting
            for _, log_row in matched.iterrows():
                row_dict = {c: log_row[c] for c in headers if c in log_row}
                log_rows.append({
                    "File":             os.path.basename(file_path),
                    "Sheet":            sn,
                    "Row Number":       int(log_row["__ws_row__"]),
                    "Condition":        cond_label,
                    "Deleted Row Data": str(row_dict),
                    "Deleted At":       now_str,
                })

            for row_idx in ws_rows:
                ws.delete_rows(row_idx)

            sheet_counts[sn] = len(ws_rows)
            total_deleted += len(ws_rows)

        wb_cache.save(wb, file_path)
        wb.close()

        # ── Write delete log ──────────────────────────────────────────────────
        log_saved_msg = ""
        if save_log and log_rows:
            try:
                if not log_path:
                    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
                    log_dir = os.path.join(os.path.dirname(file_path), "delete_logs")
                    os.makedirs(log_dir, exist_ok=True)
                    log_path = os.path.join(log_dir, f"delete_log_{ts}.xlsx")
                log_df = pd.DataFrame(log_rows)
                # Append to existing log if it already exists
                if os.path.exists(log_path):
                    existing = pd.read_excel(log_path)
                    log_df = pd.concat([existing, log_df], ignore_index=True)
                log_df.to_excel(log_path, index=False)
                log_saved_msg = f"  Log saved → {os.path.basename(log_path)}"
            except Exception as log_err:
                log_saved_msg = f"  (Log write failed: {log_err})"

        # ── Build summary ─────────────────────────────────────────────────────
        if len(target_sheets) == 1:
            summary = f"Deleted {total_deleted} row(s) from '{target_sheets[0]}'"
        else:
            per_sheet = ", ".join(f"{s}: {n}" for s, n in sheet_counts.items() if n)
            summary = (
                f"Deleted {total_deleted} row(s) across {len(target_sheets)} sheet(s)"
                + (f" ({per_sheet})" if per_sheet else "")
            )
        summary += log_saved_msg
        return True, summary

    except Exception as exc:
        return False, f"Advanced delete error: {exc}"


def _describe_filters(filters: list, combine: str) -> str:
    """Build a short human-readable description of the filter conditions."""
    if not filters:
        return "(no conditions)"
    parts = []
    for f in filters:
        col = f.get("column", "?")
        op  = f.get("operator", "eq")
        val = f.get("value", "")
        parts.append(f"[{col}] {op} '{val}'" if val else f"[{col}] {op}")
    return f" {combine} ".join(parts)

